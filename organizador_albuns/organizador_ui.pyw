# -*- coding: utf-8 -*-
"""
Janela do organizador de álbuns.

Cada programa fica numa guia:
  Buscar e copiar — planilha, pasta e plano; Simular só mostra, Copiar traz capa e lâminas.
  Renomear — troca um lote já existente pelo código SGE.
  Enviar aprovação — sobe o álbum de quem está DIAGRAMADO no MaisBook.
"""

import csv
import ctypes
import json
import os
import re
import string
import subprocess
import sys
import threading
import time
import urllib.request
from datetime import datetime
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


PASTA = Path(__file__).resolve().parent
SCRIPT = next(PASTA.glob("PROCUARA*.PY"))
SCRIPT_ENVIO = PASTA / "ENVIO_APROVACAO_MAISBOOK.PY"
CONFIG = PASTA / "organizador_config.json"
RELATORIO = PASTA / "relatorio_organizador.csv"
RELATORIO_ENVIO = PASTA / "relatorio_envio_aprovacao.csv"
PARAR_ENVIO = PASTA / "envio_aprovacao_parar.flag"
PASTA_ERROS = PASTA / "erros"
CONSOLE_PIRATA = PASTA / "console_pirata.py"
LOG_LOTE = PASTA / "organizador_saida.log"
ALTURA_CMD = 660
RESULTADOS_ENVIO_OK = {"enviado", "simulado"}

RE_RESUMO = re.compile(r"RESUMO ORGANIZADOR — (.+)")
RE_RESUMO_RENOMEIO = re.compile(
    r"RESUMO RENOMEIO — (\w+): (\d+) \| JA OK: (\d+) \| NAO ENCONTRADOS: (\d+) \| ERROS: (\d+)"
)
RE_ABA = re.compile(r"^ABA:\s*(.+)\s*$")
RE_LINK = re.compile(r"/spreadsheets/d/[\w-]+")
ORIGEM_PADRAO = r"D:\ÁLBUNS COLAÇÃO 2025.2"
PLANILHA_BAIXADA = PASTA / "planilha_organizador.xlsx"
TAMANHOS = ["20X20", "24X34", "30X20", "30X22"]
CATEGORIAS = ["CAPAS", "ESTOJO", "LAMINAS", "POSTERS"]
LETRAS = list(string.ascii_uppercase)

FUNDO = "#000000"
VERDE = "#00ff41"
VERDE_FORTE = "#b6ffb6"
VERDE_FRACO = "#1f8f3a"
BORDA = "#145214"
BOTAO = "#031403"
BOTAO_ATIVO = "#0a2a0a"
SELECAO = "#0d3d0d"
AVISO = "#e6ff4a"
ERRO = "#ff5f5f"
FONTE = ("Consolas", 10)
FONTE_NEGRITO = ("Consolas", 10, "bold")


def aplicar_tema_terminal(root):
    """Preto e verde, no mesmo espírito do terminal do Windows."""

    root.configure(background=FUNDO)
    root.option_add("*TCombobox*Listbox.background", FUNDO)
    root.option_add("*TCombobox*Listbox.foreground", VERDE)
    root.option_add("*TCombobox*Listbox.selectBackground", SELECAO)
    root.option_add("*TCombobox*Listbox.selectForeground", VERDE_FORTE)
    root.option_add("*TCombobox*Listbox.font", "Consolas 10")

    estilo = ttk.Style(root)
    try:
        estilo.theme_use("clam")
    except tk.TclError:
        pass

    estilo.configure(
        ".",
        background=FUNDO,
        foreground=VERDE,
        font=FONTE,
        troughcolor=FUNDO,
        bordercolor=BORDA,
        lightcolor=FUNDO,
        darkcolor=FUNDO,
    )
    estilo.configure("TFrame", background=FUNDO)
    estilo.configure("TLabel", background=FUNDO, foreground=VERDE, font=FONTE)
    estilo.configure("Fraco.TLabel", background=FUNDO, foreground=VERDE_FRACO, font=FONTE)
    estilo.configure(
        "TLabelframe",
        background=FUNDO,
        foreground=VERDE,
        bordercolor=BORDA,
        relief="solid",
    )
    estilo.configure("TLabelframe.Label", background=FUNDO, foreground=VERDE, font=FONTE_NEGRITO)
    estilo.configure(
        "TButton",
        background=BOTAO,
        foreground=VERDE,
        bordercolor=BORDA,
        focusthickness=1,
        focuscolor=VERDE,
        padding=(10, 4),
        font=FONTE_NEGRITO,
    )
    estilo.map(
        "TButton",
        background=[("pressed", BORDA), ("active", BOTAO_ATIVO), ("disabled", FUNDO)],
        foreground=[("disabled", VERDE_FRACO), ("pressed", VERDE_FORTE), ("active", VERDE_FORTE)],
        bordercolor=[("active", VERDE), ("focus", VERDE)],
    )
    estilo.configure("Perigo.TButton", foreground=ERRO, bordercolor="#5a1515")
    estilo.map(
        "Perigo.TButton",
        background=[("pressed", "#2a0808"), ("active", "#1a0808"), ("disabled", FUNDO)],
        foreground=[("disabled", "#5a1515"), ("active", "#ffb0b0")],
        bordercolor=[("active", ERRO), ("focus", ERRO)],
    )
    estilo.configure(
        "TEntry",
        fieldbackground=FUNDO,
        foreground=VERDE,
        insertcolor=VERDE,
        bordercolor=BORDA,
        lightcolor=BORDA,
        darkcolor=BORDA,
    )
    estilo.map(
        "TEntry",
        fieldbackground=[("disabled", FUNDO)],
        foreground=[("disabled", VERDE_FRACO)],
    )
    estilo.configure(
        "TCombobox",
        fieldbackground=FUNDO,
        background=BOTAO,
        foreground=VERDE,
        arrowcolor=VERDE,
        bordercolor=BORDA,
        insertcolor=VERDE,
        arrowsize=12,
    )
    estilo.map(
        "TCombobox",
        fieldbackground=[("readonly", FUNDO), ("disabled", FUNDO)],
        foreground=[("readonly", VERDE), ("disabled", VERDE_FRACO)],
        selectbackground=[("readonly", SELECAO)],
        selectforeground=[("readonly", VERDE_FORTE)],
        background=[("active", BOTAO_ATIVO)],
        arrowcolor=[("disabled", VERDE_FRACO)],
    )
    estilo.configure("TCheckbutton", background=FUNDO, foreground=VERDE, font=FONTE)
    estilo.map(
        "TCheckbutton",
        background=[("active", FUNDO)],
        foreground=[("active", VERDE_FORTE), ("disabled", VERDE_FRACO)],
    )
    estilo.configure("TNotebook", background=FUNDO, bordercolor=BORDA, tabmargins=(2, 4, 2, 0))
    estilo.configure(
        "TNotebook.Tab",
        background=BOTAO,
        foreground=VERDE_FRACO,
        padding=(14, 6),
        font=FONTE_NEGRITO,
        bordercolor=BORDA,
    )
    estilo.map(
        "TNotebook.Tab",
        background=[("selected", FUNDO), ("active", BOTAO_ATIVO)],
        foreground=[("selected", VERDE), ("active", VERDE_FORTE)],
        bordercolor=[("selected", VERDE)],
    )
    estilo.configure(
        "Treeview",
        background=FUNDO,
        fieldbackground=FUNDO,
        foreground=VERDE,
        bordercolor=BORDA,
        font=FONTE,
        rowheight=22,
    )
    estilo.configure(
        "Treeview.Heading",
        background=BOTAO,
        foreground=VERDE,
        font=FONTE_NEGRITO,
        bordercolor=BORDA,
        relief="flat",
    )
    estilo.map(
        "Treeview",
        background=[("selected", SELECAO)],
        foreground=[("selected", VERDE_FORTE)],
    )
    estilo.map("Treeview.Heading", background=[("active", BOTAO_ATIVO)])
    estilo.configure("TSeparator", background=BORDA)
    for nome in ("Vertical.TScrollbar", "Horizontal.TScrollbar"):
        estilo.configure(
            nome,
            background=BOTAO,
            troughcolor=FUNDO,
            arrowcolor=VERDE,
            bordercolor=FUNDO,
            relief="flat",
        )
        estilo.map(nome, background=[("active", BOTAO_ATIVO)])


def barra_escura(janela):
    try:
        janela.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(janela.winfo_id())
        valor = ctypes.c_int(1)
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            hwnd, 20, ctypes.byref(valor), ctypes.sizeof(valor)
        )
    except Exception:
        return


def pintar_console(texto):
    texto.configure(
        background=FUNDO,
        foreground=VERDE,
        insertbackground=VERDE,
        selectbackground=SELECAO,
        selectforeground=VERDE_FORTE,
        relief="flat",
        highlightthickness=1,
        highlightbackground=BORDA,
        highlightcolor=VERDE,
        bd=0,
    )
    moldura = getattr(texto, "frame", None)
    if moldura is not None:
        moldura.configure(background=FUNDO)
    barra = getattr(texto, "vbar", None)
    if barra is not None and moldura is not None:
        barra.pack_forget()
        nova = ttk.Scrollbar(moldura, command=texto.yview)
        nova.pack(side="right", fill="y")
        texto.configure(yscrollcommand=nova.set)
        texto.vbar = nova


def caminho_planilha_erros(operacao):
    PASTA_ERROS.mkdir(exist_ok=True)
    return PASTA_ERROS / f"erros_{operacao}_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx"


def planilha_erros_de_csv(relatorio, destino, eh_erro):
    """Copia do relatório CSV só as linhas com erro. Devolve quantas linhas foram gravadas."""
    with relatorio.open(encoding="utf-8-sig", newline="") as arquivo:
        leitor = csv.DictReader(arquivo, delimiter=";")
        campos = leitor.fieldnames or []
        linhas = [linha for linha in leitor if eh_erro(linha)]
    if not linhas:
        return 0
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Erros"
    aba.append(campos)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        aba.append([linha.get(campo, "") for campo in campos])
    for indice, campo in enumerate(campos, start=1):
        maior = max([len(str(campo))] + [len(str(linha.get(campo) or "")) for linha in linhas])
        aba.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 10), 80)
    aba.freeze_panes = "A2"
    livro.save(destino)
    return len(linhas)


def script_renomeio():
    for pasta in PASTA.parent.iterdir():
        if not pasta.is_dir():
            continue
        achados = list(pasta.glob("PASTA RENOMEIO.py"))
        if achados:
            return achados[0]
    return None


def baixar_planilha_google(link, destino):
    achado = re.search(r"/spreadsheets/d/([\w-]+)", link or "")
    if not achado:
        raise ValueError("Link da planilha inválido")
    gid = re.search(r"[#&?]gid=(\d+)", link)
    url = "https://docs.google.com/spreadsheets/d/" + achado.group(1) + "/export?format=xlsx"
    if gid:
        url += "&gid=" + gid.group(1)
    destino = Path(destino)
    temporario = destino.with_suffix(".baixando")
    try:
        urllib.request.urlretrieve(url, temporario)
        if temporario.read_bytes()[:2] != b"PK":
            raise ValueError(
                "o Google não devolveu a planilha "
                "(o link está compartilhado como 'qualquer pessoa com o link'?)"
            )
        temporario.replace(destino)
    except Exception:
        temporario.unlink(missing_ok=True)
        raise
    return destino


def python_console():
    """python.exe ao lado do pythonw.exe atual (a saída precisa ser capturada)."""
    exe = Path(sys.executable)
    if exe.name.lower() == "pythonw.exe":
        candidato = exe.with_name("python.exe")
        if candidato.is_file():
            return str(candidato)
    return str(exe)


def carregar_config():
    try:
        return json.loads(CONFIG.read_text(encoding="utf-8"))
    except Exception:
        return {}


class JanelaOrganizador:

    def __init__(self, root):
        self.root = root
        self.processo = None
        self.resumo = None
        self.modo = None
        self.inicio_operacao = 0.0
        self.arquivo_erros = None
        self.pediu_parar_envio = False
        self.avisos = []
        self._timer_abas = None
        self._geracao_abas = 0

        root.title("Organizador de álbuns")
        root.geometry("1000x820")
        root.minsize(860, 680)
        aplicar_tema_terminal(root)

        cfg = carregar_config()
        self.link = tk.StringVar(value=cfg.get("link", ""))
        self.planilha = tk.StringVar(value=cfg.get("planilha", ""))
        self.aba = tk.StringVar(value=cfg.get("aba", ""))
        self.origem = tk.StringVar(value=cfg.get("origem", ORIGEM_PADRAO))
        self.destino = tk.StringVar(value=cfg.get("destino", ""))
        self.col_plano = tk.StringVar(value=cfg.get("col_plano", "C"))
        self.col_nome = tk.StringVar(value=cfg.get("col_nome", "A"))
        self.col_codigo = tk.StringVar(value=cfg.get("col_codigo", "B"))
        self.plano = tk.StringVar(value=cfg.get("plano", ""))
        self.tamanhos = {
            t: tk.BooleanVar(value=t in cfg.get("tamanhos", TAMANHOS)) for t in TAMANHOS
        }
        self.categorias = {
            c: tk.BooleanVar(value=c in cfg.get("categorias", CATEGORIAS)) for c in CATEGORIAS
        }
        self.pastas_existentes = tk.StringVar(value="")
        self.envio_sge = tk.StringVar(value="")
        self.envio_limite = tk.StringVar(value="")
        self.envio_intervalo = tk.StringVar(value=str(cfg.get("envio_intervalo", 60)))
        self.envio_simular = tk.BooleanVar(value=False)
        self.envio_atualizar_planilha = tk.BooleanVar(value=cfg.get("envio_atualizar_planilha", True))
        self.envio_baixa_sge = tk.BooleanVar(value=cfg.get("envio_baixa_sge", True))
        self.status = tk.StringVar(value="Escolha a guia do programa.")
        self.combos_aba = []

        self._montar()
        self.link.trace_add("write", self._agendar_abas)
        self.planilha.trace_add("write", self._agendar_abas)
        self.destino.trace_add("write", lambda *_: self._atualizar_pastas())
        self.aba.trace_add("write", lambda *_: self._atualizar_previa())
        self.col_nome.trace_add("write", lambda *_: self._atualizar_previa())
        self.col_codigo.trace_add("write", lambda *_: self._atualizar_previa())
        self._atualizar_pastas()
        if self.link.get().strip() or self.planilha.get().strip():
            self.root.after(200, self._carregar_abas_automatico)
        barra_escura(root)

    def _montar(self):
        principal = ttk.Frame(self.root, padding=10)
        principal.pack(fill="both", expand=True)
        principal.columnconfigure(0, weight=1)
        principal.rowconfigure(2, weight=1)

        caderno = ttk.Notebook(principal)
        caderno.grid(row=0, column=0, sticky="nsew")
        self.caderno = caderno
        self._montar_busca(caderno)
        self._montar_renomeio(caderno)
        self._montar_envio(caderno)

        ttk.Label(principal, textvariable=self.status).grid(row=1, column=0, sticky="w", pady=(8, 0))
        self.texto = scrolledtext.ScrolledText(principal, height=16, state="disabled", font=FONTE)
        pintar_console(self.texto)
        self.texto.tag_configure("aviso", foreground=AVISO)
        self.texto.grid(row=2, column=0, sticky="nsew", pady=(8, 0))

    def _combo_coluna(self, pai, variavel):
        return ttk.Combobox(pai, textvariable=variavel, values=LETRAS, state="readonly", width=4)

    def _montar_busca(self, caderno):
        aba = ttk.Frame(caderno, padding=8)
        caderno.add(aba, text="Buscar e copiar")
        aba.columnconfigure(0, weight=1)

        procura = ttk.LabelFrame(aba, text="Onde procurar", padding=6)
        procura.grid(row=0, column=0, sticky="we")
        procura.columnconfigure(1, weight=1)

        ttk.Label(procura, text="Pasta").grid(row=0, column=0, sticky="w")
        ttk.Entry(procura, textvariable=self.origem).grid(row=0, column=1, sticky="we", padx=4)
        ttk.Button(procura, text="Escolher pasta...", command=self._escolher_origem).grid(row=0, column=2)
        ttk.Label(
            procura,
            text="Capas e álbuns são procurados só dentro desta pasta.",
            style="Fraco.TLabel",
        ).grid(row=1, column=1, sticky="w", padx=4, pady=(4, 0))

        dados = ttk.LabelFrame(aba, text="Planilha e destino", padding=6)
        dados.grid(row=1, column=0, sticky="we", pady=(8, 0))
        dados.columnconfigure(1, weight=1)

        ttk.Label(dados, text="Link do Google").grid(row=0, column=0, sticky="w")
        ttk.Entry(dados, textvariable=self.link).grid(row=0, column=1, columnspan=2, sticky="we", padx=4)

        ttk.Label(dados, text="Ou arquivo").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(dados, textvariable=self.planilha).grid(row=1, column=1, sticky="we", padx=4, pady=(6, 0))
        ttk.Button(dados, text="Escolher...", command=self._escolher_planilha).grid(row=1, column=2, pady=(6, 0))

        ttk.Label(dados, text="Aba").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.combo_aba = ttk.Combobox(dados, textvariable=self.aba, width=28)
        self.combo_aba.grid(row=2, column=1, sticky="w", padx=4, pady=(6, 0))
        self.combos_aba.append(self.combo_aba)
        self.botao_abas = ttk.Button(dados, text="Atualizar abas", command=self._carregar_abas)
        self.botao_abas.grid(row=2, column=2, pady=(6, 0))

        colunas = ttk.Frame(dados)
        colunas.grid(row=3, column=1, columnspan=2, sticky="w", pady=(6, 0))
        ttk.Label(dados, text="Colunas").grid(row=3, column=0, sticky="w", pady=(6, 0))
        ttk.Label(colunas, text="Nome").pack(side="left")
        self._combo_coluna(colunas, self.col_nome).pack(side="left", padx=(4, 12))
        ttk.Label(colunas, text="Código SGE").pack(side="left")
        self._combo_coluna(colunas, self.col_codigo).pack(side="left", padx=(4, 12))
        ttk.Label(colunas, text="Plano").pack(side="left")
        self.combo_col_plano = ttk.Combobox(
            colunas,
            textvariable=self.col_plano,
            values=LETRAS,
            width=4,
            state="readonly",
        )
        self.combo_col_plano.pack(side="left", padx=(4, 0))

        ttk.Label(dados, text="Destino").grid(row=4, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(dados, textvariable=self.destino).grid(row=4, column=1, sticky="we", padx=4, pady=(6, 0))
        destino_botoes = ttk.Frame(dados)
        destino_botoes.grid(row=4, column=2, pady=(6, 0))
        ttk.Button(destino_botoes, text="Escolher...", command=self._escolher_destino).pack(side="left")
        ttk.Button(destino_botoes, text="Abrir", command=self._abrir_destino).pack(side="left", padx=(4, 0))

        ttk.Label(dados, text="Filtrar plano").grid(row=5, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(dados, textvariable=self.plano, width=32).grid(row=5, column=1, sticky="w", padx=4, pady=(6, 0))
        ttk.Label(
            dados,
            text="Opcional. Ex.: Essencial, Prime, Moscou, Berlim, Diamond, 20X20",
            style="Fraco.TLabel",
        ).grid(row=6, column=1, sticky="w", padx=4)

        acoes = ttk.Frame(aba)
        acoes.grid(row=2, column=0, sticky="we", pady=(8, 0))
        self.botao_simular = ttk.Button(acoes, text="Simular", command=lambda: self._rodar(True))
        self.botao_simular.pack(side="left")
        self.botao_copiar = ttk.Button(acoes, text="Copiar", command=lambda: self._rodar(False))
        self.botao_copiar.pack(side="left", padx=(8, 0))
        self.botao_relatorio = ttk.Button(acoes, text="Abrir relatório", command=self._abrir_relatorio)
        self.botao_relatorio.pack(side="left", padx=(8, 0))

        ttk.Label(
            aba,
            text="A capa e o álbum são procurados pelo nome da pessoa. "
            "Copiar grava o lote com esse nome e, em seguida, troca pelo código SGE. "
            "A guia Renomear faz só essa troca num lote que já existe.",
            wraplength=920,
        ).grid(row=3, column=0, sticky="w", pady=(8, 0))

    def _montar_renomeio(self, caderno):
        aba = ttk.Frame(caderno, padding=8)
        caderno.add(aba, text="Renomear")
        aba.columnconfigure(0, weight=1)

        lote = ttk.LabelFrame(aba, text="Lote", padding=6)
        lote.grid(row=0, column=0, sticky="we")
        lote.columnconfigure(1, weight=1)
        ttk.Label(lote, text="Pasta do lote").grid(row=0, column=0, sticky="w")
        ttk.Entry(lote, textvariable=self.destino).grid(row=0, column=1, sticky="we", padx=4)
        ttk.Button(lote, text="Escolher...", command=self._escolher_destino).grid(row=0, column=2)
        ttk.Button(lote, text="Abrir", command=self._abrir_destino).grid(row=0, column=3, padx=(4, 0))
        ttk.Label(lote, textvariable=self.pastas_existentes, style="Fraco.TLabel").grid(
            row=1, column=1, columnspan=3, sticky="w", pady=(6, 0)
        )

        colunas = ttk.LabelFrame(aba, text="Planilha — nome e código SGE", padding=6)
        colunas.grid(row=1, column=0, sticky="we", pady=(8, 0))
        colunas.columnconfigure(6, weight=1)
        ttk.Label(colunas, text="Aba").grid(row=0, column=0, sticky="w")
        combo_aba = ttk.Combobox(colunas, textvariable=self.aba, state="readonly", width=22)
        combo_aba.grid(row=0, column=1, sticky="w", padx=(4, 12))
        self.combos_aba.append(combo_aba)
        ttk.Label(colunas, text="Coluna do nome").grid(row=0, column=2, sticky="w")
        self._combo_coluna(colunas, self.col_nome).grid(row=0, column=3, sticky="w", padx=(4, 12))
        ttk.Label(colunas, text="Coluna do código SGE").grid(row=0, column=4, sticky="w")
        self._combo_coluna(colunas, self.col_codigo).grid(row=0, column=5, sticky="w", padx=(4, 12))

        self.previa = ttk.Treeview(colunas, columns=("linha", "nome", "codigo"), show="headings", height=4)
        for col, titulo, largura in (("linha", "Linha", 60), ("nome", "Nome", 420), ("codigo", "Código SGE", 160)):
            self.previa.heading(col, text=titulo)
            self.previa.column(col, width=largura, anchor="w", stretch=(col == "nome"))
        self.previa.grid(row=1, column=0, columnspan=7, sticky="we", pady=(6, 0))

        escopo = ttk.LabelFrame(aba, text="O que processar", padding=6)
        escopo.grid(row=2, column=0, sticky="we", pady=(8, 0))
        ttk.Label(escopo, text="Tamanhos:").pack(side="left")
        for tamanho in TAMANHOS:
            ttk.Checkbutton(escopo, text=tamanho, variable=self.tamanhos[tamanho]).pack(side="left", padx=4)
        ttk.Separator(escopo, orient="vertical").pack(side="left", fill="y", padx=12)
        ttk.Label(escopo, text="Categorias:").pack(side="left")
        for categoria in CATEGORIAS:
            ttk.Checkbutton(escopo, text=categoria, variable=self.categorias[categoria]).pack(side="left", padx=4)

        acoes = ttk.Frame(aba)
        acoes.grid(row=3, column=0, sticky="we", pady=(8, 0))
        self.botao_simular_renomeio = ttk.Button(acoes, text="Simular", command=lambda: self._renomear(True))
        self.botao_simular_renomeio.pack(side="left")
        self.botao_renomear = ttk.Button(acoes, text="Renomear", command=lambda: self._renomear(False))
        self.botao_renomear.pack(side="left", padx=6)
        ttk.Label(
            aba,
            text="Usa a mesma planilha da outra guia. Simular não muda nenhum arquivo. "
            "Renomear troca capas, estojos e lâminas pelo código SGE, só nos tamanhos e categorias marcados. "
            "30X22 (e 30X22 PÓS-VENDA): capa em FRENTE e VERSO vira código_001 e código_002; sem estojo; "
            "qualquer quantidade de lâminas. POSTERS: um por aluno vira código; mais de um, código_001, código_002.",
            wraplength=920,
        ).grid(row=4, column=0, sticky="w", pady=(8, 0))

    def _montar_envio(self, caderno):
        aba = ttk.Frame(caderno, padding=8)
        caderno.add(aba, text="Enviar aprovação")
        aba.columnconfigure(0, weight=1)

        quem = ttk.LabelFrame(aba, text="Quem entra", padding=6)
        quem.grid(row=0, column=0, sticky="we")
        quem.columnconfigure(0, weight=1)
        ttk.Label(
            quem,
            text="Quem está com STATUS DIAGRAMADO na planilha ALBUNS 25.2 (Coluna C = NOME, Coluna E = SGE, Coluna K = STATUS).\n"
            "Fase 1: procura capa e álbum de todos no servidor.\n"
            "Fase 2: só quem tem os dois — sobe para ALBUM APROVAÇAO no Drive, cola o link em Álbum pronto no MaisBook, "
            "envia o e-mail de aprovação, avança a Coluna K para 'Aguardando Aprovação' e dá baixa no SGE como 'Enviado para Aprovação'.",
            wraplength=920,
        ).grid(row=0, column=0, sticky="w")

        filtros = ttk.LabelFrame(aba, text="Filtro e automação", padding=6)
        filtros.grid(row=1, column=0, sticky="we", pady=(8, 0))
        ttk.Label(filtros, text="SGE").grid(row=0, column=0, sticky="w")
        ttk.Entry(filtros, textvariable=self.envio_sge, width=14).grid(row=0, column=1, sticky="w", padx=(4, 16))
        ttk.Label(filtros, text="Limite").grid(row=0, column=2, sticky="w")
        ttk.Entry(filtros, textvariable=self.envio_limite, width=6).grid(row=0, column=3, sticky="w", padx=(4, 16))
        ttk.Label(filtros, text="Intervalo vigia (s)").grid(row=0, column=4, sticky="w")
        ttk.Entry(filtros, textvariable=self.envio_intervalo, width=6).grid(row=0, column=5, sticky="w", padx=(4, 16))
        ttk.Checkbutton(
            filtros,
            text="Simular (não sobe, não cola e não envia)",
            variable=self.envio_simular,
        ).grid(row=0, column=6, sticky="w")

        opcoes_auto = ttk.Frame(filtros)
        opcoes_auto.grid(row=1, column=0, columnspan=7, sticky="w", pady=(6, 0))
        ttk.Checkbutton(
            opcoes_auto,
            text="Atualizar Coluna K para 'Aguardando Aprovação' após envio",
            variable=self.envio_atualizar_planilha,
        ).pack(side="left", padx=(0, 16))
        ttk.Checkbutton(
            opcoes_auto,
            text="Dar baixa automática no SGE ('Enviado para Aprovação')",
            variable=self.envio_baixa_sge,
        ).pack(side="left")

        ttk.Label(
            filtros,
            text="SGE e limite são opcionais. Em 'Vigia Contínuo', o robô monitora sem fim a planilha e o servidor.",
            style="Fraco.TLabel",
        ).grid(row=2, column=0, columnspan=7, sticky="w", pady=(6, 0))

        acoes = ttk.Frame(aba)
        acoes.grid(row=2, column=0, sticky="we", pady=(8, 0))
        self.botao_envio_iniciar = ttk.Button(acoes, text="Iniciar", command=self._envio_iniciar)
        self.botao_envio_iniciar.pack(side="left")
        self.botao_envio_vigia = ttk.Button(
            acoes, text="Iniciar Vigia Contínuo", command=self._envio_vigia
        )
        self.botao_envio_vigia.pack(side="left", padx=(8, 0))
        self.botao_envio_continuar = ttk.Button(
            acoes, text="Continuar de onde parou", command=self._envio_continuar
        )
        self.botao_envio_continuar.pack(side="left", padx=(8, 0))
        self.botao_envio_fila = ttk.Button(acoes, text="Ver fila (DIAGRAMADO)", command=self._envio_listar)
        self.botao_envio_fila.pack(side="left", padx=(8, 0))
        self.botao_envio_parar = ttk.Button(
            acoes, text="Parar", command=self._envio_parar, state="disabled", style="Perigo.TButton"
        )
        self.botao_envio_parar.pack(side="left", padx=(8, 0))
        self.botao_envio_relatorio = ttk.Button(acoes, text="Abrir relatório", command=self._abrir_relatorio_envio)
        self.botao_envio_relatorio.pack(side="left", padx=(8, 0))

        ttk.Label(
            aba,
            text="Sem capa com o nome da pessoa, o envio dessa pessoa aguarda antes do Drive e do e-mail. "
            "Parar espera a pessoa atual terminar. Um segundo clique em Parar fecha na hora.",
            wraplength=920,
        ).grid(row=3, column=0, sticky="w", pady=(8, 0))

    def _escolher_planilha(self):
        caminho = filedialog.askopenfilename(
            title="Planilha",
            filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")],
        )
        if caminho:
            self.planilha.set(caminho)

    def _escolher_origem(self):
        caminho = filedialog.askdirectory(title="Pasta onde procurar capas e álbuns")
        if caminho:
            self.origem.set(caminho)

    def _escolher_destino(self):
        caminho = filedialog.askdirectory(title="Pasta de destino")
        if not caminho:
            return
        self.destino.set(os.path.normpath(caminho))
        credenciais = Path(caminho) / "CREDENCIAIS.xlsx"
        if credenciais.is_file() and not self.planilha.get().strip() and not self.link.get().strip():
            self.planilha.set(str(credenciais))
        self._atualizar_pastas()

    def _abrir_destino(self):
        pasta = self.destino.get().strip()
        if pasta and Path(pasta).is_dir():
            os.startfile(pasta)
        else:
            messagebox.showinfo("Organizador", "A pasta de destino ainda não existe.")

    def _atualizar_pastas(self):
        lote = Path(self.destino.get().strip()) if self.destino.get().strip() else None
        if not lote or not lote.is_dir():
            self.pastas_existentes.set("A pasta do lote ainda não existe.")
            return
        partes = []
        for tamanho in TAMANHOS:
            if (lote / tamanho).is_dir():
                cats = [c for c in CATEGORIAS if (lote / tamanho / c).is_dir()]
                partes.append(f"{tamanho}: {', '.join(cats) if cats else 'sem subpastas'}")
            else:
                partes.append(f"{tamanho}: não existe")
        self.pastas_existentes.set("   |   ".join(partes))

    def _caminho_planilha_local(self):
        arquivo = self.planilha.get().strip()
        if arquivo and Path(arquivo).is_file():
            return Path(arquivo)
        if self.link.get().strip() and PLANILHA_BAIXADA.is_file():
            return PLANILHA_BAIXADA
        return None

    def _atualizar_previa(self):
        if not hasattr(self, "previa"):
            return
        self.previa.delete(*self.previa.get_children())
        caminho = self._caminho_planilha_local()
        if not caminho or not self.aba.get().strip():
            return
        try:
            wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
            if self.aba.get() not in wb.sheetnames:
                wb.close()
                return
            ws = wb[self.aba.get()]
            i_nome = LETRAS.index(self.col_nome.get())
            i_cod = LETRAS.index(self.col_codigo.get())
            mostradas = 0
            total = 0
            for n, linha in enumerate(ws.iter_rows(values_only=True), start=1):
                nome = linha[i_nome] if len(linha) > i_nome else None
                cod = linha[i_cod] if len(linha) > i_cod else None
                if nome in (None, "") and cod in (None, ""):
                    continue
                total += 1
                if mostradas < 4:
                    self.previa.insert("", "end", values=(n, nome or "", cod or ""))
                    mostradas += 1
            wb.close()
            if total:
                self.status.set(f"Planilha: {total} linha(s) com nome ou código.")
        except Exception as e:
            self.status.set(f"Erro ao ler a prévia: {e}")

    def _salvar_config(self):
        dados = {
            "link": self.link.get().strip(),
            "planilha": self.planilha.get().strip(),
            "aba": self.aba.get().strip(),
            "col_plano": self.col_plano.get().strip(),
            "col_nome": self.col_nome.get().strip(),
            "col_codigo": self.col_codigo.get().strip(),
            "origem": self.origem.get().strip(),
            "destino": self.destino.get().strip(),
            "plano": self.plano.get().strip(),
            "tamanhos": [t for t, v in self.tamanhos.items() if v.get()],
            "categorias": [c for c, v in self.categorias.items() if v.get()],
            "envio_intervalo": int(self.envio_intervalo.get().strip() or 60),
            "envio_atualizar_planilha": bool(self.envio_atualizar_planilha.get()),
            "envio_baixa_sge": bool(self.envio_baixa_sge.get()),
        }
        CONFIG.write_text(json.dumps(dados, indent=2, ensure_ascii=False), encoding="utf-8")

    def _argumentos_planilha(self):
        link = self.link.get().strip()
        arquivo = self.planilha.get().strip()
        if link:
            return ["--link", link]
        if arquivo:
            return ["--planilha", arquivo]
        return None

    def _agendar_abas(self, *_):
        if self._timer_abas:
            self.root.after_cancel(self._timer_abas)
        self._timer_abas = self.root.after(700, lambda: self._carregar_abas_automatico(avisar=True))

    def _carregar_abas_automatico(self, avisar=False):
        self._timer_abas = None
        link = self.link.get().strip()
        arquivo = self.planilha.get().strip()
        if link and not RE_LINK.search(link):
            return
        if not link and not arquivo:
            return
        self._carregar_abas(avisar=avisar)

    def _carregar_abas(self, avisar=True):
        if self.processo:
            return
        args = self._argumentos_planilha()
        if not args:
            if avisar:
                messagebox.showwarning("Organizador", "Informe o link do Google ou escolha um arquivo.")
            return
        if args[0] == "--planilha" and not Path(args[1]).is_file():
            if avisar:
                messagebox.showerror("Organizador", "Arquivo não encontrado:\n\n" + args[1])
            return
        if args[0] == "--link" and not RE_LINK.search(args[1]):
            return

        self._geracao_abas += 1
        geracao = self._geracao_abas
        self.botao_abas.configure(state="disabled")
        self.status.set("Lendo as abas da planilha...")
        threading.Thread(target=self._ler_abas, args=(args, geracao, avisar), daemon=True).start()

    def _ler_abas(self, args, geracao, avisar):
        abas = []
        erro = ""
        if args[0] == "--planilha":
            try:
                livro = openpyxl.load_workbook(args[1], read_only=True, data_only=True)
                abas = list(livro.sheetnames)
                livro.close()
            except Exception as e:
                erro = str(e)
        else:
            try:
                resultado = subprocess.run(
                    [python_console(), "-u", str(SCRIPT), "--sem-pausa", "--listar-abas"] + args,
                    cwd=str(PASTA),
                    capture_output=True,
                    env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUNBUFFERED": "1"},
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                )
                saida = resultado.stdout.decode("utf-8", errors="replace")
                abas = [m.group(1) for linha in saida.splitlines() if (m := RE_ABA.match(linha))]
                if resultado.returncode != 0 and not abas:
                    erro = saida.strip() or resultado.stderr.decode("utf-8", errors="replace")
            except Exception as e:
                erro = str(e)
        self.root.after(0, self._preencher_abas, abas, erro, geracao, avisar)

    def _preencher_abas(self, abas, erro, geracao, avisar):
        if geracao != self._geracao_abas:
            return
        if not self.processo:
            self.botao_abas.configure(state="normal")
        if erro and not abas:
            self.status.set("Não consegui ler as abas.")
            if avisar:
                messagebox.showerror("Organizador", erro)
            return
        for combo in self.combos_aba:
            combo["values"] = abas
            if abas:
                combo.configure(state="readonly")
            else:
                combo.configure(state="normal")
        if abas and self.aba.get() not in abas:
            self.aba.set(abas[0])
        self.status.set(f"{len(abas)} aba(s) encontrada(s)." if abas else "Nenhuma aba encontrada.")
        self._atualizar_previa()

    def _rodar(self, simular):
        if self.processo:
            return
        args_planilha = self._argumentos_planilha()
        origem = self.origem.get().strip()
        destino = self.destino.get().strip()
        if not args_planilha:
            messagebox.showwarning("Organizador", "Informe o link do Google ou escolha um arquivo.")
            return
        if not origem:
            messagebox.showwarning("Organizador", "Escolha a pasta onde procurar as capas e os álbuns.")
            return
        if not destino:
            messagebox.showwarning("Organizador", "Escolha a pasta de destino. Se ela não existir, será criada no modo Copiar.")
            return
        if not SCRIPT.is_file():
            messagebox.showerror("Organizador", "Script não encontrado:\n\n" + str(SCRIPT))
            return
        if not CONSOLE_PIRATA.is_file():
            messagebox.showerror("Organizador", "Script não encontrado:\n\n" + str(CONSOLE_PIRATA))
            return
        if not simular and not messagebox.askyesno(
            "Copiar",
            "Os arquivos serão COPIADOS para:\n\n"
            + destino
            + "\n\nSe essa pasta não existir, ela será criada com a organização\n"
            "tamanho / CAPAS, LAMINAS/{nome da pessoa} e ESTOJO (24X34 e 30X20).\n"
            "Depois da cópia, capa e lâminas são renomeadas com o código SGE.\n\n"
            "Os originais do servidor são mantidos. Continuar?",
        ):
            return

        self._salvar_config()
        args = args_planilha + [
            "--origem", origem,
            "--destino", destino,
            "--sem-pausa",
            "--simular" if simular else "--real",
        ]
        if self.aba.get().strip():
            args += ["--aba", self.aba.get().strip()]
        if self.col_plano.get().strip():
            args += ["--col-plano", self.col_plano.get().strip()]
        if self.col_nome.get().strip():
            args += ["--col-nome", self.col_nome.get().strip()]
        if self.col_codigo.get().strip():
            args += ["--col-codigo", self.col_codigo.get().strip()]
        if self.plano.get().strip():
            args += ["--plano", self.plano.get().strip()]

        self.resumo = None
        self.modo = "busca"
        self.inicio_operacao = time.time()
        self._limpar()
        self._escrever("> " + SCRIPT.name + " " + " ".join(f'"{a}"' if " " in a else a for a in args) + "\n")
        self._habilitar(False)
        self.status.set("Simulando..." if simular else "Copiando...")

        try:
            LOG_LOTE.unlink(missing_ok=True)
        except OSError:
            pass
        x, y, largura, altura = self._posicao_cmd()
        try:
            self.processo = subprocess.Popen(
                [
                    python_console(), "-u", str(CONSOLE_PIRATA),
                    "--log", str(LOG_LOTE),
                    "--x", str(x), "--y", str(y),
                    "--largura", str(largura), "--altura", str(altura),
                    "--titulo", "Separando o lote - " + ("Simular" if simular else "Copiar"),
                    "--", python_console(), "-u", str(SCRIPT),
                ] + args,
                cwd=str(PASTA),
                creationflags=getattr(subprocess, "CREATE_NEW_CONSOLE", 0),
            )
        except Exception as e:
            self.processo = None
            self._habilitar(True)
            self.status.set("Não foi possível iniciar o script.")
            self._escrever(str(e) + "\n")
            return

        threading.Thread(target=self._acompanhar_log, args=(self.processo,), daemon=True).start()

    def _posicao_cmd(self):
        self.root.update_idletasks()
        x = self.root.winfo_x()
        y = self.root.winfo_rooty() + self.root.winfo_height() + 8
        largura = self.root.winfo_width() + 16
        altura = ALTURA_CMD
        tela = self.root.winfo_screenheight()
        if y + altura > tela - 40:
            y = max(tela - 40 - altura, 0)
        return x, y, largura, altura

    def _acompanhar_log(self, processo):
        codigo = None
        arquivo = None
        sobra = ""
        try:
            while codigo is None:
                if arquivo is None:
                    if LOG_LOTE.is_file():
                        arquivo = open(LOG_LOTE, "r", encoding="utf-8", errors="replace")
                    elif processo.poll() is not None:
                        break
                    else:
                        time.sleep(0.1)
                        continue
                pedaco = arquivo.read()
                if not pedaco:
                    if processo.poll() is not None:
                        break
                    time.sleep(0.1)
                    continue
                sobra += pedaco
                linhas = sobra.split("\n")
                sobra = linhas.pop()
                texto = []
                for linha in linhas:
                    if linha.startswith("__FIM__"):
                        partes = linha.split()
                        codigo = int(partes[1]) if len(partes) > 1 and partes[1].lstrip("-").isdigit() else 1
                        break
                    texto.append(linha + "\n")
                if texto:
                    self.root.after(0, self._escrever, "".join(texto))
        finally:
            if arquivo:
                arquivo.close()
        if codigo is None:
            codigo = processo.poll()
            if codigo is None or codigo == 0:
                codigo = 1
            self.root.after(0, self._escrever, "\n⚠ O CMD foi fechado antes do fim do processo.\n")
        self.root.after(0, self._finalizar, codigo)

    def _ler_saida(self):
        processo = self.processo
        for bruto in iter(processo.stdout.readline, b""):
            self.root.after(0, self._escrever, bruto.decode("utf-8", errors="replace"))
        codigo = processo.wait()
        self.root.after(0, self._finalizar, codigo)

    def _gerar_planilha_erros(self, modo):
        """Caminho e quantidade de linhas da planilha de erros desta operação, ou (None, 0)."""
        try:
            if modo == "renomeio":
                arquivo = self.arquivo_erros
                if arquivo and arquivo.is_file():
                    livro = openpyxl.load_workbook(arquivo, read_only=True)
                    linhas = max(livro.active.max_row - 1, 0)
                    livro.close()
                    return arquivo, linhas
                return None, 0
            if modo == "busca":
                relatorio, operacao = RELATORIO, "organizador"
                eh_erro = lambda linha: (linha.get("STATUS") or "").strip().upper() != "OK"
            elif modo == "envio":
                relatorio, operacao = RELATORIO_ENVIO, "envio"
                eh_erro = lambda linha: (linha.get("RESULTADO") or "").strip().lower() not in RESULTADOS_ENVIO_OK
            else:
                return None, 0
            if not relatorio.is_file() or relatorio.stat().st_mtime < self.inicio_operacao:
                return None, 0
            destino = caminho_planilha_erros(operacao)
            linhas = planilha_erros_de_csv(relatorio, destino, eh_erro)
            return (destino, linhas) if linhas else (None, 0)
        except Exception as e:
            self._escrever(f"\n⚠ Não consegui gerar a planilha de erros: {e}\n")
            return None, 0

    def _finalizar(self, codigo):
        avisos = list(self.avisos)
        self.avisos = []
        self.processo = None
        planilha_erros, linhas_erros = self._gerar_planilha_erros(self.modo)
        self.arquivo_erros = None
        self.modo = None
        PARAR_ENVIO.unlink(missing_ok=True)
        self._habilitar(True)
        self._atualizar_pastas()
        self._atualizar_previa()
        if avisos:
            self.status.set(f"{len(avisos)} aviso(s): itens não encontrados ou pessoas duplicadas.")
        elif codigo != 0:
            self.status.set(f"Terminou com erro (código {codigo}). Veja a saída.")
        elif self.resumo:
            self.status.set("Concluído: " + self.resumo)
        else:
            self.status.set("Concluído.")
        if planilha_erros:
            self._escrever(f"\n>>> Planilha de erros ({linhas_erros} linha(s)): {planilha_erros}\n")
            self.status.set(self.status.get() + f"  |  Planilha de erros: {linhas_erros} linha(s)")
        if avisos:
            texto = "\n".join(item.replace("NÃO ENCONTREI: ", "Não encontrei: ").replace("DUPLICADO: ", "Duplicado: ") for item in avisos[:25])
            if len(avisos) > 25:
                texto += f"\n\n... e mais {len(avisos) - 25}."
            if planilha_erros:
                texto += f"\n\nPlanilha de erros gerada ({linhas_erros} linha(s)). Abrir agora?"
                if messagebox.askyesno("Itens não encontrados", texto, icon="warning"):
                    os.startfile(planilha_erros)
                return
            messagebox.showwarning("Itens não encontrados", texto)
        if planilha_erros and messagebox.askyesno(
            "Planilha de erros",
            f"A operação terminou com {linhas_erros} erro(s).\n\n{planilha_erros}\n\nAbrir agora?",
            icon="warning",
        ):
            os.startfile(planilha_erros)

    def _abrir_relatorio(self):
        if not RELATORIO.is_file():
            messagebox.showinfo("Organizador", "Ainda não há relatório. Rode Simular ou Copiar primeiro.")
            return
        os.startfile(RELATORIO)

    def _abrir_relatorio_envio(self):
        if not RELATORIO_ENVIO.is_file():
            messagebox.showinfo("Enviar aprovação", "Ainda não há relatório. Rode o envio primeiro.")
            return
        os.startfile(RELATORIO_ENVIO)

    def _habilitar(self, habilitar):
        estado = "normal" if habilitar else "disabled"
        self.botao_simular.configure(state=estado)
        self.botao_copiar.configure(state=estado)
        self.botao_abas.configure(state=estado)
        self.botao_simular_renomeio.configure(state=estado)
        self.botao_renomear.configure(state=estado)
        self.botao_envio_iniciar.configure(state=estado)
        self.botao_envio_vigia.configure(state=estado)
        self.botao_envio_continuar.configure(state=estado)
        self.botao_envio_fila.configure(state=estado)
        self.botao_envio_relatorio.configure(state=estado)
        if habilitar or self.modo != "envio":
            self.pediu_parar_envio = False
            self.botao_envio_parar.configure(state="disabled", text="Parar")
        else:
            self.botao_envio_parar.configure(state="normal", text="Parar")

    def _limpar(self):
        self.avisos = []
        self.texto.configure(state="normal")
        self.texto.delete("1.0", "end")
        self.texto.configure(state="disabled")

    def _escrever(self, texto):
        achado = RE_RESUMO.search(texto)
        if achado:
            self.resumo = achado.group(1)
        renomeio = RE_RESUMO_RENOMEIO.search(texto)
        if renomeio:
            self.resumo = (
                f"{renomeio.group(1).capitalize()} {renomeio.group(2)} | Já ok {renomeio.group(3)} | "
                f"Não encontrados {renomeio.group(4)} | Erros {renomeio.group(5)}"
            )
        aviso = None
        for linha in texto.splitlines():
            if linha.strip().startswith(("NÃO ENCONTREI:", "DUPLICADO:")):
                aviso = linha.strip()
                self.avisos.append(aviso)
        if aviso:
            self.status.set(aviso)
        self.texto.configure(state="normal")
        if aviso:
            self.texto.insert("end", texto, "aviso")
        else:
            self.texto.insert("end", texto)
        self.texto.see("end")
        self.texto.configure(state="disabled")


    def _renomear(self, simular):
        if self.processo:
            return
        lote = self.destino.get().strip()
        tamanhos = [t for t, v in self.tamanhos.items() if v.get()]
        categorias = [c for c, v in self.categorias.items() if v.get()]
        script = script_renomeio()
        link = self.link.get().strip()
        arquivo = self.planilha.get().strip()

        if not lote or not Path(lote).is_dir():
            messagebox.showwarning("Renomear", "Escolha uma pasta de lote que já exista.")
            return
        if not tamanhos or not categorias:
            messagebox.showwarning("Renomear", "Marque pelo menos um tamanho e uma categoria.")
            return
        if script is None or not script.is_file():
            messagebox.showerror("Renomear", "Não encontrei o PASTA RENOMEIO.py.")
            return
        if link and not RE_LINK.search(link):
            messagebox.showwarning("Renomear", "O link da planilha não parece um Google Sheets.")
            return
        if not link and (not arquivo or not Path(arquivo).is_file()):
            messagebox.showwarning("Renomear", "Informe o link do Google ou escolha a planilha.")
            return
        if not simular and not messagebox.askyesno(
            "Renomear",
            "Os arquivos e pastas do lote serão RENOMEADOS de verdade:\n\n"
            + lote
            + "\n\nRecomendado simular antes. Continuar?",
        ):
            return

        self._salvar_config()
        self.resumo = None
        self.modo = "renomeio"
        self._limpar()
        self._habilitar(False)
        self.status.set("Baixando a planilha..." if link else ("Simulando renomeio..." if simular else "Renomeando..."))
        threading.Thread(
            target=self._preparar_renomeio,
            args=(simular, script, lote, tamanhos, categorias),
            daemon=True,
        ).start()

    def _preparar_renomeio(self, simular, script, lote, tamanhos, categorias):
        try:
            link = self.link.get().strip()
            if link:
                planilha = baixar_planilha_google(link, PLANILHA_BAIXADA)
            else:
                planilha = Path(self.planilha.get().strip())
        except Exception as e:
            self.root.after(0, self._falha_inicio, str(e))
            return
        self.root.after(0, self._comecar_renomeio, simular, script, lote, planilha, tamanhos, categorias)

    def _falha_inicio(self, mensagem):
        self.processo = None
        self._habilitar(True)
        self.status.set("Não foi possível iniciar.")
        self._escrever(mensagem + "\n")

    def _comecar_renomeio(self, simular, script, lote, planilha, tamanhos, categorias):
        args = [
            "--lote", lote,
            "--planilha", str(planilha),
            "--col-nome", self.col_nome.get(),
            "--col-codigo", self.col_codigo.get(),
            "--tamanhos", ",".join(tamanhos),
            "--categorias", ",".join(categorias),
            "--sem-pausa",
            "--simular" if simular else "--real",
        ]
        if self.aba.get().strip():
            args += ["--aba", self.aba.get().strip()]
        self.inicio_operacao = time.time()
        self.arquivo_erros = caminho_planilha_erros("renomeio_simulado" if simular else "renomeio")
        args += ["--erros", str(self.arquivo_erros)]

        self._escrever("> " + script.name + " " + " ".join(f'"{a}"' if " " in a else a for a in args) + "\n")
        self.status.set("Simulando renomeio..." if simular else "Renomeando...")
        ambiente = dict(os.environ)
        ambiente["PYTHONIOENCODING"] = "utf-8"
        ambiente["PYTHONUNBUFFERED"] = "1"
        try:
            self.processo = subprocess.Popen(
                [python_console(), "-u", str(script)] + args,
                cwd=str(script.parent),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                stdin=subprocess.DEVNULL,
                env=ambiente,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except Exception as e:
            self._falha_inicio(str(e))
            return
        threading.Thread(target=self._ler_saida, daemon=True).start()

    def _envio_argumentos(self):
        argumentos = []
        sge = self.envio_sge.get().strip()
        if sge:
            if not re.fullmatch(r"\d+", sge.split(".")[-1]):
                messagebox.showerror("Enviar aprovação", "Digite só o número do SGE.")
                return None
            argumentos += ["--sge", sge]
        limite = self.envio_limite.get().strip()
        if limite:
            if not limite.isdigit() or int(limite) < 1:
                messagebox.showerror("Enviar aprovação", "O limite é a quantidade de pessoas.")
                return None
            argumentos += ["--limite", limite]
        if not self.envio_atualizar_planilha.get():
            argumentos += ["--sem-planilha"]
        if not self.envio_baixa_sge.get():
            argumentos += ["--sem-sge"]
        return argumentos

    def _envio_iniciar(self):
        if self.processo:
            return
        argumentos = self._envio_argumentos()
        if argumentos is None:
            return
        simular = self.envio_simular.get()
        if simular:
            argumentos = argumentos + ["--simular"]
        elif not messagebox.askyesno(
            "Enviar de verdade",
            "Isso sobe as fotos no Drive, grava o link em Álbum pronto, envia o e-mail, "
            "atualiza a Coluna K para 'Aguardando Aprovação' e dá baixa no SGE.\n\nContinuar?",
        ):
            return
        self._envio_executar(
            argumentos,
            "Simulando o envio..." if simular else "Enviando álbum para aprovação...",
        )

    def _envio_vigia(self):
        if self.processo:
            return
        argumentos = self._envio_argumentos()
        if argumentos is None:
            return
        intervalo = self.envio_intervalo.get().strip() or "60"
        if not intervalo.isdigit() or int(intervalo) < 10:
            messagebox.showerror("Vigia Contínuo", "O intervalo do vigia deve ser de pelo menos 10 segundos.")
            return
        argumentos += ["--vigia", "--intervalo", intervalo]
        simular = self.envio_simular.get()
        if simular:
            argumentos += ["--simular"]
        self._envio_executar(
            argumentos,
            f"Vigia Contínuo ativo (a cada {intervalo}s) — monitorando DIAGRAMADO...",
        )

    def _envio_continuar(self):
        if self.processo:
            return
        argumentos = self._envio_argumentos()
        if argumentos is None:
            return
        argumentos = argumentos + ["--continuar"]
        simular = self.envio_simular.get()
        if simular:
            argumentos = argumentos + ["--simular"]
        elif not messagebox.askyesno(
            "Continuar o envio",
            "Continua de quem ainda não foi enviado.\nSobe fotos, grava o link, manda o e-mail e atualiza Coluna K/SGE.\n\nContinuar?",
        ):
            return
        self._envio_executar(
            argumentos,
            "Simulando o envio..." if simular else "Continuando o envio...",
        )

    def _envio_listar(self):
        if self.processo:
            return
        self._envio_executar(["--listar"], "Listando quem está DIAGRAMADO...")

    def _envio_executar(self, argumentos, status):
        if not SCRIPT_ENVIO.is_file():
            messagebox.showerror("Enviar aprovação", "Script não encontrado:\n\n" + str(SCRIPT_ENVIO))
            return
        PARAR_ENVIO.unlink(missing_ok=True)
        self.pediu_parar_envio = False
        self.modo = "envio"
        self.inicio_operacao = time.time()
        self.resumo = None
        self._limpar()
        self._escrever("> " + SCRIPT_ENVIO.name + " " + " ".join(argumentos) + "\n")
        self._habilitar(False)
        self.status.set(status)
        ambiente = dict(os.environ)
        ambiente["PYTHONIOENCODING"] = "utf-8"
        ambiente["PYTHONUNBUFFERED"] = "1"
        try:
            self.processo = subprocess.Popen(
                [python_console(), "-u", str(SCRIPT_ENVIO)] + argumentos,
                cwd=str(PASTA),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                stdin=subprocess.DEVNULL,
                env=ambiente,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except Exception as e:
            self.processo = None
            self.modo = None
            self._habilitar(True)
            self.status.set("Não foi possível iniciar o envio.")
            self._escrever(str(e) + "\n")
            return
        threading.Thread(target=self._ler_saida, daemon=True).start()

    def _envio_parar(self):
        if not self.processo or self.modo != "envio":
            return
        if not self.pediu_parar_envio:
            PARAR_ENVIO.write_text("parar", encoding="utf-8")
            self.pediu_parar_envio = True
            self.botao_envio_parar.configure(text="Forçar parada")
            self.status.set("Parando ao terminar a pessoa atual...")
            self._escrever("\n>>> Pedido para parar ao terminar a pessoa atual.\n")
            return
        if not messagebox.askyesno(
            "Forçar parada",
            "Fechar o envio agora? A pessoa no meio não fica marcada como enviada.",
        ):
            return
        subprocess.run(
            ["taskkill", "/PID", str(self.processo.pid), "/T", "/F"],
            capture_output=True,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )


def main():
    root = tk.Tk()
    JanelaOrganizador(root)
    root.mainloop()


if __name__ == "__main__":
    main()
