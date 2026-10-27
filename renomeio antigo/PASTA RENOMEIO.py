import argparse
import os
import re
import sys
import unicodedata
from openpyxl import load_workbook


# ==========================================================
# CONFIGURAÇÃO
# ==========================================================

# Pasta principal onde estão:
#
# LOTE 11
# ├── 20X20
# │   ├── CAPAS
# │   └── LAMINAS
# │
# ├── 24X34
# │   ├── CAPAS
# │   ├── ESTOJO
# │   └── LAMINAS
# │
# └── 30X20
#     ├── CAPAS
#     ├── ESTOJO
#     └── LAMINAS
#
PASTA_LOTE = r"C:\Users\REIS\Documents\PAPA 2\LOTE 14"


# ==========================================================
# PLANILHA
# ==========================================================

PLANILHA = r"C:\Users\REIS\Documents\PAPA 2\LOTE 14\CREDENCIAIS.xlsx"

# Nome da aba
ABA = "Página1"


# ==========================================================
# COLUNAS DA PLANILHA
# ==========================================================

# Coluna A = Nome do cliente
COLUNA_NOME = 1

# Coluna B = Código SGE — É DESTA COLUNA QUE O CÓDIGO FINAL É OBTIDO
COLUNA_CODIGO = 2


# ==========================================================
# MODO TESTE
# ==========================================================
#
# True  = NÃO RENOMEIA NADA
# False = RENOMEIA DE VERDADE
#
MODO_TESTE = False 


# ==========================================================
# O QUE PROCESSAR
# ==========================================================

TAMANHOS = ["20X20", "24X34", "30X20", "30X22"]

CATEGORIAS = ["CAPAS", "ESTOJO", "LAMINAS", "POSTERS"]

# Sem estojo; a capa vem em dois arquivos: FRENTE -> código_001, VERSO -> código_002.
TAMANHOS_FRENTE_VERSO = {"30X22"}

# True = espera ENTER no fim (uso com duplo clique)
PAUSAR_NO_FIM = True


# ==========================================================
# CONTADORES DO RESUMO
# ==========================================================

CONTADORES = {
    "renomeados": 0,
    "ja_ok": 0,
    "nao_encontrados": 0,
    "erros": 0,
}

NAO_ENCONTRADOS = []

# Linhas da planilha de erros: {"TIPO", "LOCAL", "MOTIVO"}.
ERROS = []

# --erros: onde gravar a planilha de erros (.xlsx).
ARQUIVO_ERROS = None


def contar(tipo, nome=None, motivo=""):
    CONTADORES[tipo] += 1
    if tipo == "nao_encontrados" and nome:
        NAO_ENCONTRADOS.append(nome)
    if tipo == "nao_encontrados":
        ERROS.append({
            "TIPO": "NÃO ENCONTRADO",
            "LOCAL": nome or "",
            "MOTIVO": motivo or "Nome não encontrado na planilha",
        })
    elif tipo == "erros":
        ERROS.append({"TIPO": "ERRO", "LOCAL": nome or "", "MOTIVO": motivo})


def salvar_planilha_erros(caminho):

    from openpyxl import Workbook
    from openpyxl.styles import Font

    livro = Workbook()
    aba = livro.active
    aba.title = "Erros"
    aba.append(["TIPO", "LOCAL", "MOTIVO"])

    for celula in aba[1]:
        celula.font = Font(bold=True)

    for erro in ERROS:
        aba.append([erro["TIPO"], erro["LOCAL"], erro["MOTIVO"]])

    aba.column_dimensions["A"].width = 18
    aba.column_dimensions["B"].width = 70
    aba.column_dimensions["C"].width = 60
    aba.freeze_panes = "A2"

    try:
        livro.save(caminho)
        print()
        print(f"PLANILHA DE ERROS: {caminho} ({len(ERROS)} linha(s))")
    except PermissionError:
        print()
        print(f"⚠ Não consegui gravar {caminho} — ela está aberta no Excel?")


def pausar():
    if PAUSAR_NO_FIM:
        input("Pressione ENTER para fechar...")


# ==========================================================
# NORMALIZAR TEXTO
# ==========================================================

def normalizar(texto):

    if texto is None:
        return ""

    texto = str(texto).strip().lower()

    # Remove acentos
    texto = unicodedata.normalize(
        "NFD",
        texto
    )

    texto = "".join(
        caractere
        for caractere in texto
        if unicodedata.category(caractere) != "Mn"
    )

    # Troca caracteres especiais por espaço
    texto = re.sub(
        r"[^a-z0-9]+",
        " ",
        texto
    )

    # Remove espaços duplicados
    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


# ==========================================================
# IDENTIFICAR NOME DA PESSOA NO ARQUIVO / PASTA
# ==========================================================

def nome_pessoa_do_arquivo(nome_arquivo):

    nome = os.path.splitext(nome_arquivo)[0].strip()

    # Remove tamanho e lado da capa no final, em qualquer ordem.
    # Ex.: "Maria Silva FRENTE 30X22" -> "Maria Silva"
    anterior = None

    while anterior != nome:

        anterior = nome

        nome = re.sub(
            r"[\s_-]*(20X20|24X34|30X20|30X22|FRENTE|VERSO|POSTERS?|\(\s*\d+\s*\))\s*$",
            "",
            nome,
            flags=re.IGNORECASE
        )

    # Remove numeração de lâmina no final, se existir.
    # Ex.: "João da Silva_001" -> "João da Silva"
    # O separador é obrigatório: "123456" é um código SGE, não "123" + "456".
    nome = re.sub(
        r"(?<=\S)[\s_-]+\d{3}\s*$",
        "",
        nome
    )

    return nome.strip()


def eh_arquivo_de_sistema(nome_arquivo):

    """Thumbs.db e afins não são lâmina nem capa: não recebem número."""

    nome = nome_arquivo.lower()

    return nome in {"thumbs.db", "ehthumbs.db", "desktop.ini", ".ds_store"} or nome.startswith("._")


def lado_da_capa(nome_arquivo):

    """1 = FRENTE, 2 = VERSO, None = sem lado no nome."""

    palavras = normalizar(
        os.path.splitext(nome_arquivo)[0]
    ).split()

    if "frente" in palavras:
        return 1

    if "verso" in palavras:
        return 2

    return None


def codigo_ja_aplicado(nome, clientes):

    """Código da planilha se o nome já for 'código', 'código_001' etc. Qualquer quantidade de dígitos."""

    base = os.path.splitext(nome)[0].strip()

    encontrado = re.fullmatch(
        r"(\d+)(?:[\s_-]+\d{3})?",
        base
    )

    if not encontrado:
        return None

    codigo = encontrado.group(1)

    if codigo in set(clientes.values()):
        return codigo

    return None


def encontrar_cliente_por_nome(nome_arquivo, clientes):

    nome_pessoa = nome_pessoa_do_arquivo(nome_arquivo)
    nome_normalizado = normalizar(nome_pessoa)

    # Busca exata pelo nome da planilha.
    if nome_normalizado in clientes:
        return clientes[nome_normalizado]

    # Busca parcial somente quando houver um único candidato.
    candidatos = []

    for nome_planilha, codigo in clientes.items():
        if (
            nome_normalizado in nome_planilha
            or nome_planilha in nome_normalizado
        ):
            candidatos.append(codigo)

    if len(candidatos) == 1:
        return candidatos[0]

    return None


# ==========================================================
# PEGAR CÓDIGO
# ==========================================================

def pegar_codigo(codigo):

    if codigo is None:
        return None

    # O código SGE deve vir EXATAMENTE da coluna B da planilha.
    # Não corta o primeiro dígito e não acrescenta zero no final.
    if isinstance(codigo, int):
        return str(codigo)

    if isinstance(codigo, float):
        # Excel pode entregar 53220 como 53220.0.
        # Remove somente a parte decimal .0.
        if codigo.is_integer():
            return str(int(codigo))
        return str(codigo).strip()

    codigo = str(codigo).strip()

    # Se estiver como texto, preserva os dígitos exatamente.
    # Remove somente .0 no final, caso exista.
    codigo = re.sub(r"\.0+$", "", codigo)

    numeros = re.sub(r"\D", "", codigo)

    if not numeros:
        return None

    return numeros


# ==========================================================
# CARREGAR CLIENTES DA PLANILHA
# ==========================================================

def carregar_clientes():

    print()
    print("=" * 70)
    print("LENDO PLANILHA")
    print("=" * 70)

    # Verificar se a planilha existe
    if not os.path.exists(PLANILHA):

        print()
        print("❌ ERRO: PLANILHA NÃO ENCONTRADA")
        print()

        print(
            "Caminho informado:"
        )

        print(
            PLANILHA
        )

        return {}


    # Abrir planilha
    try:

        workbook = load_workbook(
            PLANILHA,
            data_only=True
        )

    except Exception as erro:

        print()
        print(
            "❌ ERRO AO ABRIR A PLANILHA"
        )

        print(
            erro
        )

        return {}


    # Verificar aba
    if ABA not in workbook.sheetnames:

        print()
        print(
            f"❌ ERRO: A aba "
            f"'{ABA}' não existe."
        )

        print()
        print(
            "Abas disponíveis:"
        )

        for nome_aba in workbook.sheetnames:

            print(
                f" - {nome_aba}"
            )

        return {}


    sheet = workbook[ABA]

    clientes = {}


    # ======================================================
    # LER LINHAS DA PLANILHA
    # ======================================================

    for linha in sheet.iter_rows(
        min_row=1,
        values_only=True
    ):

        # Verificar se existem colunas suficientes
        if len(linha) < max(
            COLUNA_NOME,
            COLUNA_CODIGO
        ):

            continue


        nome = linha[
            COLUNA_NOME - 1
        ]

        codigo = linha[
            COLUNA_CODIGO - 1
        ]


        # Ignorar linhas vazias
        if not nome or not codigo:
            continue


        nome = str(
            nome
        ).strip()


        codigo = pegar_codigo(
            codigo
        )


        if not codigo:
            continue


        nome_normalizado = normalizar(
            nome
        )


        if not nome_normalizado:
            continue


        clientes[
            nome_normalizado
        ] = codigo


    print()
    print(
        f"✓ Clientes carregados: "
        f"{len(clientes)}"
    )

    return clientes


# ==========================================================
# PROCURAR CLIENTE NA PLANILHA
# ==========================================================

def procurar_cliente(
    nome,
    clientes
):

    nome_normalizado = normalizar(
        nome
    )


    # ======================================================
    # PRIMEIRO: BUSCA EXATA
    # ======================================================

    if nome_normalizado in clientes:

        return clientes[
            nome_normalizado
        ]


    # ======================================================
    # SEGUNDO: BUSCA POR PARTE DO NOME
    # ======================================================

    candidatos = []


    for nome_planilha, codigo in clientes.items():

        if (
            nome_normalizado in nome_planilha
            or nome_planilha in nome_normalizado
        ):

            candidatos.append(
                (
                    nome_planilha,
                    codigo
                )
            )


    # Só aceita se houver exatamente um candidato
    if len(candidatos) == 1:

        return candidatos[0][1]


    return None


# ==========================================================
# PROCESSAR CAPAS / ESTOJO
# ==========================================================

def processar_arquivo_individual(
    arquivo,
    clientes,
    tamanho,
    categoria
):

    nome_original = arquivo.name

    print()
    print("-" * 70)
    print(f"[{tamanho} / {categoria}]")
    print(f"Arquivo atual: {nome_original}")

    if codigo_ja_aplicado(nome_original, clientes):
        print("✓ Já está com o código SGE da planilha.")
        contar("ja_ok")
        return True

    # ======================================================
    # IDENTIFICAR A PESSOA PELO NOME ATUAL
    # ======================================================

    nome_pessoa = nome_pessoa_do_arquivo(nome_original)

    codigo = encontrar_cliente_por_nome(
        nome_original,
        clientes
    )

    print(f"Pessoa identificada: {nome_pessoa}")

    if not codigo:
        print("❌ NOME NÃO ENCONTRADO NA PLANILHA")
        print("   O arquivo NÃO será renomeado.")
        contar("nao_encontrados", f"{tamanho}/{categoria}: {nome_original}")
        return False

    # ======================================================
    # O CÓDIGO VEM EXCLUSIVAMENTE DA PLANILHA
    # ======================================================

    print(f"✓ Código SGE encontrado na planilha: {codigo}")

    extensao = os.path.splitext(nome_original)[1]

    # Resultado: SOMENTE o código SGE da coluna B.
    return aplicar_novo_nome(arquivo, f"{codigo}{extensao}")


def aplicar_novo_nome(arquivo, novo_nome):

    nome_original = arquivo.name
    pasta = os.path.dirname(arquivo.path)
    caminho_original = arquivo.path
    novo_caminho = os.path.join(pasta, novo_nome)

    if nome_original.lower() == novo_nome.lower():
        print("✓ Já está com o nome correto.")
        contar("ja_ok")
        return True

    if os.path.exists(novo_caminho):
        print("⚠ NÃO SERIA RENOMEADO")
        print(f"Já existe: {novo_nome}")
        contar("erros", arquivo.path, f"Já existe um arquivo chamado {novo_nome}")
        return False

    if MODO_TESTE:
        print("🔵 MODO TESTE")
        print(f"  {nome_original}")
        print("       ↓")
        print(f"  {novo_nome}")
        print("  ⚠ Nenhuma alteração foi feita.")
        contar("renomeados")
        return True

    try:
        os.rename(
            caminho_original,
            novo_caminho
        )

        print("✅ RENOMEADO")
        print(f"  {nome_original} → {novo_nome}")

        contar("renomeados")
        return True

    except Exception as erro:
        print("❌ ERRO AO RENOMEAR:")
        print(erro)
        contar("erros", arquivo.path, f"Erro ao renomear para {novo_nome}: {erro}")
        return False


# ==========================================================
# CAPAS EM FRENTE E VERSO (30X22)
# ==========================================================

def processar_capas_frente_verso(
    pasta_categoria,
    clientes,
    tamanho
):

    try:
        arquivos = [
            item for item in os.scandir(pasta_categoria)
            if item.is_file()
        ]
    except Exception as erro:
        print("❌ ERRO AO ACESSAR CAPAS")
        print(erro)
        contar("erros", pasta_categoria, f"Erro ao abrir a pasta: {erro}")
        return

    for arquivo in arquivos:

        nome_original = arquivo.name

        if eh_arquivo_de_sistema(nome_original):
            continue

        print()
        print("-" * 70)
        print(f"[{tamanho} / CAPAS — FRENTE/VERSO]")
        print(f"Arquivo atual: {nome_original}")

        if codigo_ja_aplicado(nome_original, clientes):
            print("✓ Já está com o código SGE da planilha.")
            contar("ja_ok")
            continue

        lado = lado_da_capa(nome_original)

        if lado is None:
            print("❌ SEM FRENTE/VERSO NO NOME")
            print("   O arquivo NÃO será renomeado.")
            contar("erros", arquivo.path, "Capa sem FRENTE ou VERSO no nome")
            continue

        nome_pessoa = nome_pessoa_do_arquivo(nome_original)
        codigo = encontrar_cliente_por_nome(nome_original, clientes)

        print(f"Pessoa identificada: {nome_pessoa}")
        print(f"Lado: {'FRENTE' if lado == 1 else 'VERSO'}")

        if not codigo:
            print("❌ NOME NÃO ENCONTRADO NA PLANILHA")
            print("   O arquivo NÃO será renomeado.")
            contar("nao_encontrados", f"{tamanho}/CAPAS: {nome_original}")
            continue

        print(f"✓ Código SGE encontrado na planilha: {codigo}")

        extensao = os.path.splitext(nome_original)[1]

        aplicar_novo_nome(arquivo, f"{codigo}_{lado:03d}{extensao}")


# ==========================================================
# POSTERS
# ==========================================================

def processar_posters(
    pasta_categoria,
    clientes,
    tamanho
):

    """Um poster por aluno vira código.ext; dois ou mais viram código_001, código_002..."""

    try:
        arquivos = [
            item for item in os.scandir(pasta_categoria)
            if item.is_file() and not eh_arquivo_de_sistema(item.name)
        ]
    except Exception as erro:
        print("❌ ERRO AO ACESSAR POSTERS")
        print(erro)
        contar("erros", pasta_categoria, f"Erro ao abrir a pasta: {erro}")
        return

    por_codigo = {}

    for arquivo in ordenar_arquivos_lamina(arquivos):

        codigo = codigo_ja_aplicado(
            arquivo.name,
            clientes
        ) or encontrar_cliente_por_nome(
            arquivo.name,
            clientes
        )

        if not codigo:
            print()
            print("-" * 70)
            print(f"[{tamanho} / POSTERS]")
            print(f"Arquivo atual: {arquivo.name}")
            print(f"Pessoa identificada: {nome_pessoa_do_arquivo(arquivo.name)}")
            print("❌ NOME NÃO ENCONTRADO NA PLANILHA")
            print("   O arquivo NÃO será renomeado.")
            contar("nao_encontrados", f"{tamanho}/POSTERS: {arquivo.name}")
            continue

        por_codigo.setdefault(codigo, []).append(arquivo)

    for codigo, grupo in por_codigo.items():

        print()
        print("-" * 70)
        print(f"[{tamanho} / POSTERS] Código SGE: {codigo} ({len(grupo)} arquivo(s))")

        # Os já renomeados primeiro, na ordem do número; depois os novos.
        grupo.sort(key=lambda arquivo: (codigo_ja_aplicado(arquivo.name, clientes) is None,))

        destinos = []

        for numero, arquivo in enumerate(grupo, start=1):

            extensao = os.path.splitext(arquivo.name)[1]

            if len(grupo) == 1:
                destinos.append(f"{codigo}{extensao}")
            else:
                destinos.append(f"{codigo}_{numero:03d}{extensao}")

        if all(
            arquivo.name.lower() == destino.lower()
            for arquivo, destino in zip(grupo, destinos)
        ):
            print("✓ Já está com o código SGE da planilha.")
            contar("ja_ok")
            continue

        nomes_do_grupo = {arquivo.name.lower() for arquivo in grupo}
        ocupados = [
            destino for destino in destinos
            if destino.lower() not in nomes_do_grupo
            and os.path.exists(os.path.join(pasta_categoria, destino))
        ]

        if ocupados:
            print("⚠ NÃO RENOMEADO")
            print(f"Já existe: {', '.join(ocupados)}")
            contar(
                "erros",
                ", ".join(arquivo.path for arquivo in grupo),
                f"Já existe: {', '.join(ocupados)}",
            )
            continue

        for arquivo, destino in zip(grupo, destinos):
            print(f"  {arquivo.name}")
            print("       ↓")
            print(f"  {destino}")

        if MODO_TESTE:
            print("🔵 MODO TESTE")
            print("  ⚠ Nenhuma alteração foi feita.")
            contar("renomeados")
            continue

        temporarios = []

        try:
            for numero, arquivo in enumerate(grupo, start=1):
                temporario = os.path.join(
                    pasta_categoria,
                    f"__RENOMEANDO_POSTER_{numero:05d}__{os.path.splitext(arquivo.name)[1]}"
                )
                os.rename(arquivo.path, temporario)
                temporarios.append(temporario)

            for temporario, destino in zip(temporarios, destinos):
                os.rename(temporario, os.path.join(pasta_categoria, destino))

            print("✅ RENOMEADO")
            contar("renomeados")

        except Exception as erro:
            print("❌ ERRO AO RENOMEAR:")
            print(erro)
            contar("erros", pasta_categoria, f"Erro ao renomear os posters do código {codigo}: {erro}")


# ==========================================================
# ORDEM DAS LÂMINAS
# ==========================================================

def base_e_numero_lamina(nome):

    """'Foto (10)' vira ('Foto', 10). Sem (n), o número fica vazio."""

    stem = os.path.splitext(nome)[0].strip()

    encontrado = re.search(
        r"^(.*?)(?:\s*\(\s*(\d+)\s*\))\s*$",
        stem
    )

    if not encontrado or encontrado.group(2) is None:

        return stem, None

    return encontrado.group(1).strip(), int(encontrado.group(2))


def ordenar_arquivos_lamina(arquivos):

    """(10) não pode vir antes de (2). O arquivo sem (n) é a página 1."""

    bases = set()

    for arquivo in arquivos:

        base, numero = base_e_numero_lamina(arquivo.name)

        if numero is not None:

            bases.add(normalizar(base))

    def chave(arquivo):

        base, numero = base_e_numero_lamina(arquivo.name)
        nome = arquivo.name.lower()

        if numero is not None:

            return (numero, 1, nome)

        # Cópia original do Windows: "Foto.jpg" ao lado de "Foto (2).jpg".
        if bases and normalizar(base) in bases:

            return (1, 0, nome)

        return (10**9, 1, nome)

    return sorted(arquivos, key=chave)


# ==========================================================
# PROCESSAR LAMINAS
# ==========================================================

def processar_laminas(
    pasta_laminas,
    clientes,
    tamanho
):

    print()
    print()
    print("=" * 70)
    print(f"{tamanho} / LAMINAS")
    print("=" * 70)

    try:
        pastas_clientes = [
            item for item in os.scandir(pasta_laminas)
            if item.is_dir()
        ]
    except Exception as erro:
        print("❌ ERRO AO ACESSAR LAMINAS")
        print(erro)
        contar("erros", pasta_laminas, f"Erro ao abrir a pasta: {erro}")
        return

    for pasta_cliente in pastas_clientes:

        nome_pasta_atual = pasta_cliente.name

        # ==================================================
        # IDENTIFICAR A PESSOA PELO NOME DA PASTA
        # ==================================================

        # A pasta ainda tem o nome do cliente; o código vem da coluna B.
        # Pasta já renomeada: o nome é o próprio código, de qualquer tamanho.
        codigo = codigo_ja_aplicado(
            nome_pasta_atual,
            clientes
        ) or encontrar_cliente_por_nome(
            nome_pasta_atual,
            clientes
        )

        print()
        print("-" * 70)
        print("PASTA:")
        print(f"  {nome_pasta_atual}")

        if not codigo:
            print("❌ NOME NÃO ENCONTRADO NA PLANILHA")
            print("   Esta pasta NÃO será renomeada.")
            contar("nao_encontrados", f"{tamanho}/LAMINAS: {nome_pasta_atual}")
            continue

        print(f"✓ Código SGE encontrado na planilha: {codigo}")

        # ==================================================
        # PEGAR ARQUIVOS
        # ==================================================

        try:
            arquivos = [
                arquivo for arquivo in os.scandir(pasta_cliente.path)
                if arquivo.is_file() and not eh_arquivo_de_sistema(arquivo.name)
            ]
        except Exception as erro:
            print("❌ ERRO AO ACESSAR A PASTA")
            print(erro)
            contar("erros", pasta_cliente.path, f"Erro ao abrir a pasta: {erro}")
            continue

        if not arquivos:
            print("Nenhum arquivo encontrado.")

        # (10) em ordem alfabética vinha antes de (2) e da foto sem número.
        arquivos = ordenar_arquivos_lamina(arquivos)

        if arquivos and nome_pasta_atual == codigo and all(
            arquivo.name.lower()
            == f"{codigo}_{numero:03d}{os.path.splitext(arquivo.name)[1]}".lower()
            for numero, arquivo in enumerate(arquivos, start=1)
        ):
            print(f"✓ Pasta e lâminas já estão com o código {codigo}")
            contar("ja_ok")
            continue

        # ==================================================
        # MODO TESTE
        # ==================================================

        if MODO_TESTE:

            print("🔵 SIMULAÇÃO DAS LÂMINAS:")

            for numero, arquivo in enumerate(
                arquivos,
                start=1
            ):

                extensao = os.path.splitext(
                    arquivo.name
                )[1]

                novo_nome = (
                    f"{codigo}_{numero:03d}{extensao}"
                )

                print()
                print(f"  {arquivo.name}")
                print("       ↓")
                print(f"  {novo_nome}")

            print()
            print("📁 SIMULAÇÃO DA PASTA:")
            print(f"  {nome_pasta_atual}")
            print("       ↓")
            print(f"  {codigo}")

            print("  ⚠ Nenhum arquivo ou pasta foi alterado.")
            contar("renomeados")
            continue

        # ==================================================
        # RENOMEAÇÃO TEMPORÁRIA
        # ==================================================

        temporarios = []
        erro_temporario = False
        motivo_temporario = ""

        for numero, arquivo in enumerate(
            arquivos,
            start=1
        ):

            extensao = os.path.splitext(
                arquivo.name
            )[1]

            nome_temporario = (
                f"__RENOMEANDO_LAMINA_{numero:05d}__"
                f"{extensao}"
            )

            caminho_temporario = os.path.join(
                pasta_cliente.path,
                nome_temporario
            )

            if os.path.exists(caminho_temporario):
                print("❌ ARQUIVO TEMPORÁRIO JÁ EXISTE:")
                print(nome_temporario)
                motivo_temporario = f"Arquivo temporário já existe: {nome_temporario}"
                erro_temporario = True
                break

            try:
                os.rename(
                    arquivo.path,
                    caminho_temporario
                )

                temporarios.append(
                    (
                        caminho_temporario,
                        arquivo.name,
                        extensao
                    )
                )

            except Exception as erro:
                print("❌ ERRO AO CRIAR NOME TEMPORÁRIO:")
                print(arquivo.name)
                print(erro)
                motivo_temporario = f"Erro ao preparar {arquivo.name}: {erro}"
                erro_temporario = True
                break

        if erro_temporario:
            contar(
                "erros",
                pasta_cliente.path,
                motivo_temporario,
            )
            continue

        # ==================================================
        # RENOMEAR ARQUIVOS COM O CÓDIGO SGE
        # ==================================================

        sucesso_arquivos = True
        motivos_arquivos = []

        for numero, dados in enumerate(
            temporarios,
            start=1
        ):

            caminho_temporario = dados[0]
            nome_original = dados[1]
            extensao = dados[2]

            novo_nome = (
                f"{codigo}_{numero:03d}{extensao}"
            )

            novo_caminho = os.path.join(
                pasta_cliente.path,
                novo_nome
            )

            if os.path.exists(novo_caminho):
                print("⚠ NÃO RENOMEADO:")
                print(f"Já existe: {novo_nome}")
                motivos_arquivos.append(f"{nome_original}: já existe {novo_nome}")
                sucesso_arquivos = False
                continue

            try:
                os.rename(
                    caminho_temporario,
                    novo_caminho
                )

                print()
                print(f"✓ {nome_original} → {novo_nome}")

            except Exception as erro:
                print("❌ ERRO:")
                print(erro)
                motivos_arquivos.append(f"{nome_original}: {erro}")
                sucesso_arquivos = False

        # ==================================================
        # RENOMEAR A PASTA COM O CÓDIGO SGE
        # ==================================================

        if sucesso_arquivos:

            caminho_pasta_nova = os.path.join(
                pasta_laminas,
                codigo
            )

            if nome_pasta_atual == codigo:

                print()
                print(f"✓ A pasta já está com o nome {codigo}")
                contar("ja_ok")

            elif os.path.exists(caminho_pasta_nova):

                print()
                print("⚠ PASTA NÃO RENOMEADA")
                print(f"Já existe uma pasta chamada: {codigo}")
                contar("erros", pasta_cliente.path, f"Já existe uma pasta chamada {codigo}")

            else:

                try:
                    os.rename(
                        pasta_cliente.path,
                        caminho_pasta_nova
                    )

                    print()
                    print("📁 PASTA RENOMEADA")
                    print(
                        f"  {nome_pasta_atual} → {codigo}"
                    )
                    contar("renomeados")

                except Exception as erro:
                    print("❌ ERRO AO RENOMEAR A PASTA:")
                    print(erro)
                    contar("erros", pasta_cliente.path, f"Erro ao renomear a pasta para {codigo}: {erro}")

        else:

            print()
            print("⚠ A PASTA NÃO SERÁ RENOMEADA.")
            print("Houve problema em um ou mais arquivos.")
            contar(
                "erros",
                pasta_cliente.path,
                "Pasta não renomeada; problema nas lâminas: " + "; ".join(motivos_arquivos),
            )


# ==========================================================
# PROGRAMA PRINCIPAL
# ==========================================================

def coluna_para_numero(valor):
    """Aceita letra (A, B, AA) ou número (1, 2)."""

    valor = str(valor).strip().upper()

    if valor.isdigit():
        return int(valor)

    numero = 0
    for letra in valor:
        if not "A" <= letra <= "Z":
            raise argparse.ArgumentTypeError(f"Coluna inválida: {valor}")
        numero = numero * 26 + (ord(letra) - ord("A") + 1)
    return numero


def lista_maiuscula(texto):
    return [p.strip().upper() for p in str(texto).split(",") if p.strip()]


def aplicar_argumentos(argv=None):

    global PASTA_LOTE, PLANILHA, ABA, COLUNA_NOME, COLUNA_CODIGO
    global MODO_TESTE, TAMANHOS, CATEGORIAS, PAUSAR_NO_FIM, ARQUIVO_ERROS

    parser = argparse.ArgumentParser(
        description="Renomeia capas, estojos e lâminas pelo código SGE da planilha."
    )
    parser.add_argument("--lote", help="Pasta do lote (com 20X20, 24X34, 30X20)")
    parser.add_argument("--planilha", help="Planilha (padrão: CREDENCIAIS.xlsx dentro do lote)")
    parser.add_argument("--aba", help="Nome da aba")
    parser.add_argument("--col-nome", type=coluna_para_numero, help="Coluna do nome (A ou 1)")
    parser.add_argument("--col-codigo", type=coluna_para_numero, help="Coluna do código SGE (B ou 2)")
    parser.add_argument("--tamanhos", type=lista_maiuscula, help="Ex.: 20X20,24X34,30X20")
    parser.add_argument("--categorias", type=lista_maiuscula, help="Ex.: CAPAS,ESTOJO,LAMINAS")
    parser.add_argument("--simular", action="store_true", help="Não renomeia nada")
    parser.add_argument("--real", action="store_true", help="Renomeia de verdade")
    parser.add_argument("--sem-pausa", action="store_true", help="Não espera ENTER no fim")
    parser.add_argument("--erros", help="Planilha .xlsx onde gravar os erros no fim")

    args = parser.parse_args(argv)

    if args.erros:
        ARQUIVO_ERROS = args.erros

    if args.lote:
        PASTA_LOTE = args.lote
        if not args.planilha:
            PLANILHA = os.path.join(PASTA_LOTE, "CREDENCIAIS.xlsx")

    if args.planilha:
        PLANILHA = args.planilha

    if args.aba:
        ABA = args.aba

    if args.col_nome:
        COLUNA_NOME = args.col_nome

    if args.col_codigo:
        COLUNA_CODIGO = args.col_codigo

    if args.tamanhos:
        TAMANHOS = args.tamanhos

    if args.categorias:
        CATEGORIAS = args.categorias

    if args.simular:
        MODO_TESTE = True
    elif args.real:
        MODO_TESTE = False

    if args.sem_pausa:
        PAUSAR_NO_FIM = False


# ==========================================================
# PASTAS DE CADA TAMANHO (INCLUI PÓS-VENDA)
# ==========================================================

def eh_pos_venda(nome):

    return "posvenda" in normalizar(nome).replace(" ", "")


def pastas_do_tamanho(tamanho):

    """[(rótulo, caminho)]. No 30X22 entram também '30X22 PÓS-VENDA' e '30X22/PÓS-VENDA'."""

    pastas = []
    principal = os.path.join(PASTA_LOTE, tamanho)

    if os.path.isdir(principal):
        pastas.append((tamanho, principal))

    if tamanho not in TAMANHOS_FRENTE_VERSO:
        return pastas

    prefixo = normalizar(tamanho)

    try:
        for item in os.scandir(PASTA_LOTE):
            if (
                item.is_dir()
                and item.name.upper() != tamanho
                and normalizar(item.name).startswith(prefixo)
                and eh_pos_venda(item.name)
            ):
                pastas.append((item.name, item.path))
    except Exception as erro:
        print(f"❌ ERRO AO LER O LOTE: {erro}")

    if os.path.isdir(principal):
        try:
            for item in os.scandir(principal):
                if item.is_dir() and eh_pos_venda(item.name):
                    pastas.append((f"{tamanho}/{item.name}", item.path))
        except Exception as erro:
            print(f"❌ ERRO AO LER {tamanho}: {erro}")

    return pastas


def processar_pasta_tamanho(
    pasta_tamanho,
    rotulo,
    tamanho,
    categorias,
    clientes
):

    print()
    print()
    print("#" * 70)
    print(f"PROCESSANDO {rotulo}")
    print("#" * 70)

    frente_verso = tamanho in TAMANHOS_FRENTE_VERSO

    for categoria in categorias:

        if frente_verso and categoria == "ESTOJO":

            print()
            print(f"↳ {rotulo} / ESTOJO: este tamanho não tem estojo. Ignorando.")
            continue

        pasta_categoria = os.path.join(
            pasta_tamanho,
            categoria
        )

        # Exemplo: 20X20 não possui ESTOJO.
        if not os.path.isdir(pasta_categoria):

            print()
            print(f"↳ {rotulo} / {categoria}: não existe. Ignorando.")
            continue

        if categoria == "LAMINAS":

            processar_laminas(
                pasta_categoria,
                clientes,
                rotulo
            )

            continue

        if categoria == "POSTERS":

            print()
            print()
            print("=" * 70)
            print(f"{rotulo} / POSTERS")
            print("=" * 70)

            processar_posters(
                pasta_categoria,
                clientes,
                rotulo
            )

            continue

        if frente_verso and categoria == "CAPAS":

            print()
            print()
            print("=" * 70)
            print(f"{rotulo} / CAPAS (FRENTE = _001, VERSO = _002)")
            print("=" * 70)

            processar_capas_frente_verso(
                pasta_categoria,
                clientes,
                rotulo
            )

            continue

        # CAPAS / ESTOJO
        print()
        print()
        print("=" * 70)
        print(f"{rotulo} / {categoria}")
        print("=" * 70)

        try:
            arquivos = [
                item for item in os.scandir(pasta_categoria)
                if item.is_file()
            ]
        except Exception as erro:
            print()
            print("❌ ERRO AO ACESSAR:")
            print(erro)
            contar("erros", pasta_categoria, f"Erro ao abrir a pasta: {erro}")
            continue

        for arquivo in arquivos:

            processar_arquivo_individual(
                arquivo,
                clientes,
                rotulo,
                categoria
            )


def main(argv=None):

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    aplicar_argumentos(argv)

    print()
    print("=" * 70)

    print(
        "AUTOMAÇÃO DE RENOMEAÇÃO"
    )

    print(
        "=" * 70
    )


    # ======================================================
    # MOSTRAR MODO
    # ======================================================

    print()


    if MODO_TESTE:

        print(
            "🔵 MODO TESTE ATIVADO"
        )

        print()

        print(
            "NENHUM ARQUIVO OU PASTA "
            "SERÁ RENOMEADO."
        )


    else:

        print(
            "🔴 MODO REAL ATIVADO"
        )

        print()

        print(
            "OS ARQUIVOS E PASTAS "
            "SERÃO RENOMEADOS."
        )


    print()


    # ======================================================
    # VERIFICAR PASTA DO LOTE
    # ======================================================

    if not os.path.isdir(
        PASTA_LOTE
    ):

        print(
            "❌ PASTA DO LOTE NÃO ENCONTRADA"
        )

        print()

        print(
            PASTA_LOTE
        )

        pausar()

        return


    print(
        "Pasta do lote:"
    )

    print(
        PASTA_LOTE
    )


    # ======================================================
    # CARREGAR PLANILHA
    # ======================================================

    clientes = carregar_clientes()


    if not clientes:

        print()

        print(
            "❌ Nenhum cliente foi carregado."
        )

        pausar()

        return


    tamanhos = TAMANHOS

    categorias = CATEGORIAS

    print()
    print("Tamanhos:", ", ".join(tamanhos))
    print("Categorias:", ", ".join(categorias))


    # ======================================================
    # PROCESSAR TUDO
    # ======================================================

    for tamanho in tamanhos:

        pastas = pastas_do_tamanho(tamanho)


        # --------------------------------------------------
        # TAMANHO NÃO EXISTE
        # --------------------------------------------------

        if not pastas:

            print()

            print(
                f"⚠ Pasta não encontrada:"
                f" {tamanho}"
            )

            continue

        for rotulo, pasta_tamanho in pastas:

            processar_pasta_tamanho(
                pasta_tamanho,
                rotulo,
                tamanho,
                categorias,
                clientes
            )


    # ======================================================
    # FINAL
    # ======================================================

    print()
    print()

    print(
        "=" * 70
    )

    print(
        "SIMULAÇÃO FINALIZADA"
        if MODO_TESTE
        else "PROCESSO FINALIZADO"
    )

    print(
        "=" * 70
    )


    if MODO_TESTE:

        print()

        print(
            "🔵 MODO TESTE:"
        )

        print(
            "NENHUM ARQUIVO OU PASTA "
            "FOI ALTERADO."
        )

        print()

        print(
            "Confira principalmente:"
        )

        print(
            "1. Nome da pessoa identificado no arquivo/pasta"
        )

        print(
            "2. Código SGE encontrado na coluna B da planilha"
        )

        print(
            "3. Arquivos renomeados pelo código SGE da planilha"
        )

        print(
            "4. Numeração das lâminas"
        )

        print(
            "5. Pastas das lâminas renomeadas pelo código SGE da planilha"
        )


    else:

        print()

        print(
            "✅ PROCESSO FINALIZADO."
        )


    if NAO_ENCONTRADOS:

        print()
        print("NÃO ENCONTRADOS NA PLANILHA:")

        for item in NAO_ENCONTRADOS:
            print(f"  - {item}")

    print()
    print("=" * 70)
    print(
        "RESUMO RENOMEIO — "
        f"{'SIMULADOS' if MODO_TESTE else 'RENOMEADOS'}: {CONTADORES['renomeados']} | "
        f"JA OK: {CONTADORES['ja_ok']} | "
        f"NAO ENCONTRADOS: {CONTADORES['nao_encontrados']} | "
        f"ERROS: {CONTADORES['erros']}"
    )
    print("=" * 70)

    if ARQUIVO_ERROS and ERROS:
        salvar_planilha_erros(ARQUIVO_ERROS)
    elif ARQUIVO_ERROS:
        print()
        print("Nenhum erro: planilha de erros não foi gerada.")

    print()

    pausar()


# ==========================================================
# INICIAR
# ==========================================================

if __name__ == "__main__":

    main()