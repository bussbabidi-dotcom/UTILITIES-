# -*- coding: utf-8 -*-
"""
Janela do renomeio de capas, estojos e lâminas pelo código SGE.

Escolha a pasta do lote e a planilha, confira as colunas e clique em
Simular ou Renomear. Roda "PASTA RENOMEIO.py" em processo separado.
"""

import json
import os
import re
import string
import subprocess
import sys
import threading
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl


PASTA = Path(__file__).resolve().parent
SCRIPT = PASTA / "PASTA RENOMEIO.py"
CONFIG = PASTA / "renomeio_config.json"

TAMANHOS = ["20X20", "24X34", "30X20"]
CATEGORIAS = ["CAPAS", "ESTOJO", "LAMINAS"]
LETRAS = list(string.ascii_uppercase)

RE_RESUMO = re.compile(
    r"RESUMO RENOMEIO — (\w+): (\d+) \| JA OK: (\d+) \| NAO ENCONTRADOS: (\d+) \| ERROS: (\d+)"
)


def python_console():
    """python.exe ao lado do pythonw.exe atual (a saída precisa ser capturada)."""
    exe = Path(sys.executable)
    if exe.name.lower() == "pythonw.exe":
        candidato = exe.with_name("python.exe")
        if candidato.is_file():
            return str(candidato)
    return str(exe)


def carregar_config():
    try:
        return json.loads(CONFIG.read_text(encoding="utf-8"))
    except Exception:
        return {}


class JanelaRenomeio:

    def __init__(self, root):
        self.root = root
        self.processo = None
        self.resumo = None

        root.title("Renomeio pelo código SGE")
        root.geometry("1000x760")
        root.minsize(820, 600)

        cfg = carregar_config()

        self.lote = tk.StringVar(value=cfg.get("lote", ""))
        self.planilha = tk.StringVar(value=cfg.get("planilha", ""))
        self.aba = tk.StringVar(value=cfg.get("aba", ""))
        self.col_nome = tk.StringVar(value=cfg.get("col_nome", "A"))
        self.col_codigo = tk.StringVar(value=cfg.get("col_codigo", "B"))
        self.tamanhos = {t: tk.BooleanVar(value=t in cfg.get("tamanhos", TAMANHOS)) for t in TAMANHOS}
        self.categorias = {c: tk.BooleanVar(value=c in cfg.get("categorias", CATEGORIAS)) for c in CATEGORIAS}
        self.status = tk.StringVar(value="Escolha a pasta do lote.")
        self.pastas_existentes = tk.StringVar(value="")

        self._montar()

        if self.planilha.get():
            self._carregar_abas(manter_aba=True)
        self._atualizar_pastas()

    # --------------------------------------------------------
    # Montagem
    # --------------------------------------------------------

    def _montar(self):
        principal = ttk.Frame(self.root, padding=10)
        principal.pack(fill="both", expand=True)
        principal.columnconfigure(0, weight=1)
        principal.rowconfigure(4, weight=1)

        # ---------------- Arquivos ----------------
        arquivos = ttk.LabelFrame(principal, text="Lote e planilha", padding=6)
        arquivos.grid(row=0, column=0, sticky="we")
        arquivos.columnconfigure(1, weight=1)

        ttk.Label(arquivos, text="Pasta do lote").grid(row=0, column=0, sticky="w")
        ttk.Entry(arquivos, textvariable=self.lote).grid(row=0, column=1, sticky="we", padx=4)
        ttk.Button(arquivos, text="Escolher...", command=self._escolher_lote).grid(row=0, column=2)
        ttk.Button(arquivos, text="Abrir", command=self._abrir_lote).grid(row=0, column=3, padx=(4, 0))

        ttk.Label(arquivos, text="Planilha").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(arquivos, textvariable=self.planilha).grid(row=1, column=1, sticky="we", padx=4, pady=(6, 0))
        ttk.Button(arquivos, text="Escolher...", command=self._escolher_planilha).grid(row=1, column=2, pady=(6, 0))

        ttk.Label(arquivos, textvariable=self.pastas_existentes, foreground="#555").grid(
            row=2, column=1, columnspan=3, sticky="w", pady=(6, 0)
        )

        # ---------------- Colunas + prévia ----------------
        colunas = ttk.LabelFrame(principal, text="Planilha", padding=6)
        colunas.grid(row=1, column=0, sticky="we", pady=(8, 0))
        colunas.columnconfigure(6, weight=1)

        ttk.Label(colunas, text="Aba").grid(row=0, column=0, sticky="w")
        self.combo_aba = ttk.Combobox(colunas, textvariable=self.aba, state="readonly", width=22)
        self.combo_aba.grid(row=0, column=1, sticky="w", padx=(4, 12))
        self.combo_aba.bind("<<ComboboxSelected>>", lambda e: self._atualizar_previa())

        ttk.Label(colunas, text="Coluna do nome").grid(row=0, column=2, sticky="w")
        c1 = ttk.Combobox(colunas, textvariable=self.col_nome, values=LETRAS, state="readonly", width=4)
        c1.grid(row=0, column=3, sticky="w", padx=(4, 12))
        c1.bind("<<ComboboxSelected>>", lambda e: self._atualizar_previa())

        ttk.Label(colunas, text="Coluna do código SGE").grid(row=0, column=4, sticky="w")
        c2 = ttk.Combobox(colunas, textvariable=self.col_codigo, values=LETRAS, state="readonly", width=4)
        c2.grid(row=0, column=5, sticky="w", padx=(4, 12))
        c2.bind("<<ComboboxSelected>>", lambda e: self._atualizar_previa())

        self.previa = ttk.Treeview(colunas, columns=("linha", "nome", "codigo"), show="headings", height=4)
        for col, titulo, largura in (("linha", "Linha", 60), ("nome", "Nome", 420), ("codigo", "Código SGE", 160)):
            self.previa.heading(col, text=titulo)
            self.previa.column(col, width=largura, anchor="w", stretch=(col == "nome"))
        self.previa.grid(row=1, column=0, columnspan=7, sticky="we", pady=(6, 0))

        # ---------------- O que processar ----------------
        escopo = ttk.LabelFrame(principal, text="O que processar", padding=6)
        escopo.grid(row=2, column=0, sticky="we", pady=(8, 0))

        ttk.Label(escopo, text="Tamanhos:").pack(side="left")
        for t in TAMANHOS:
            ttk.Checkbutton(escopo, text=t, variable=self.tamanhos[t]).pack(side="left", padx=4)

        ttk.Separator(escopo, orient="vertical").pack(side="left", fill="y", padx=12)

        ttk.Label(escopo, text="Categorias:").pack(side="left")
        for c in CATEGORIAS:
            ttk.Checkbutton(escopo, text=c, variable=self.categorias[c]).pack(side="left", padx=4)

        # ---------------- Ações ----------------
        acoes = ttk.Frame(principal)
        acoes.grid(row=3, column=0, sticky="we", pady=(8, 0))

        self.botao_simular = ttk.Button(acoes, text="Simular", command=lambda: self.executar(True))
        self.botao_simular.pack(side="left")

        self.botao_renomear = ttk.Button(acoes, text="Renomear", command=lambda: self.executar(False))
        self.botao_renomear.pack(side="left", padx=6)

        ttk.Label(acoes, textvariable=self.status).pack(side="left", padx=12)

        # ---------------- Saída ----------------
        saida = ttk.LabelFrame(principal, text="Saída", padding=6)
        saida.grid(row=4, column=0, sticky="nsew", pady=(8, 0))
        saida.columnconfigure(0, weight=1)
        saida.rowconfigure(0, weight=1)

        self.texto_saida = tk.Text(saida, wrap="word", state="disabled",
                                   background="#1e1e1e", foreground="#dddddd")
        self.texto_saida.grid(row=0, column=0, sticky="nsew")
        barra = ttk.Scrollbar(saida, orient="vertical", command=self.texto_saida.yview)
        barra.grid(row=0, column=1, sticky="ns")
        self.texto_saida.configure(yscrollcommand=barra.set)

    # --------------------------------------------------------
    # Lote e planilha
    # --------------------------------------------------------

    def _escolher_lote(self):
        inicial = self.lote.get() or str(PASTA.parent)
        caminho = filedialog.askdirectory(title="Pasta do lote", initialdir=inicial)
        if not caminho:
            return
        self.lote.set(os.path.normpath(caminho))

        credenciais = Path(caminho) / "CREDENCIAIS.xlsx"
        if credenciais.is_file():
            self.planilha.set(str(credenciais))
            self._carregar_abas()
        self._atualizar_pastas()

    def _abrir_lote(self):
        if self.lote.get() and Path(self.lote.get()).is_dir():
            os.startfile(self.lote.get())

    def _escolher_planilha(self):
        inicial = self.lote.get() or str(PASTA)
        caminho = filedialog.askopenfilename(
            title="Planilha",
            initialdir=inicial,
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")],
        )
        if caminho:
            self.planilha.set(os.path.normpath(caminho))
            self._carregar_abas()

    def _atualizar_pastas(self):
        lote = Path(self.lote.get()) if self.lote.get() else None
        if not lote or not lote.is_dir():
            self.pastas_existentes.set("")
            return
        partes = []
        for t in TAMANHOS:
            if (lote / t).is_dir():
                cats = [c for c in CATEGORIAS if (lote / t / c).is_dir()]
                partes.append(f"{t}: {', '.join(cats) if cats else 'sem subpastas'}")
            else:
                partes.append(f"{t}: não existe")
        self.pastas_existentes.set("   |   ".join(partes))

    def _carregar_abas(self, manter_aba=False):
        caminho = self.planilha.get()
        self.combo_aba["values"] = []
        if not caminho or not Path(caminho).is_file():
            self.previa.delete(*self.previa.get_children())
            return
        try:
            wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
            abas = wb.sheetnames
            wb.close()
        except Exception as e:
            self.status.set(f"Não foi possível abrir a planilha: {e}")
            return
        self.combo_aba["values"] = abas
        if not (manter_aba and self.aba.get() in abas):
            self.aba.set(abas[0] if abas else "")
        self._atualizar_previa()

    def _atualizar_previa(self):
        self.previa.delete(*self.previa.get_children())
        caminho = self.planilha.get()
        if not caminho or not Path(caminho).is_file() or not self.aba.get():
            return
        try:
            wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
            ws = wb[self.aba.get()]
            i_nome = LETRAS.index(self.col_nome.get())
            i_cod = LETRAS.index(self.col_codigo.get())
            mostradas = 0
            total = 0
            for n, linha in enumerate(ws.iter_rows(values_only=True), start=1):
                nome = linha[i_nome] if len(linha) > i_nome else None
                cod = linha[i_cod] if len(linha) > i_cod else None
                if nome in (None, "") and cod in (None, ""):
                    continue
                total += 1
                if mostradas < 4:
                    self.previa.insert("", "end", values=(n, nome or "", cod or ""))
                    mostradas += 1
            wb.close()
            self.status.set(f"Planilha: {total} linha(s) com dados.")
        except Exception as e:
            self.status.set(f"Erro ao ler a planilha: {e}")

    # --------------------------------------------------------
    # Execução
    # --------------------------------------------------------

    def _salvar_config(self):
        dados = {
            "lote": self.lote.get(),
            "planilha": self.planilha.get(),
            "aba": self.aba.get(),
            "col_nome": self.col_nome.get(),
            "col_codigo": self.col_codigo.get(),
            "tamanhos": [t for t, v in self.tamanhos.items() if v.get()],
            "categorias": [c for c, v in self.categorias.items() if v.get()],
        }
        try:
            CONFIG.write_text(json.dumps(dados, indent=2, ensure_ascii=False), encoding="utf-8")
        except Exception:
            pass

    def executar(self, simular):
        if self.processo is not None:
            return

        lote = self.lote.get().strip()
        planilha = self.planilha.get().strip()
        tamanhos = [t for t, v in self.tamanhos.items() if v.get()]
        categorias = [c for c, v in self.categorias.items() if v.get()]

        if not lote or not Path(lote).is_dir():
            messagebox.showwarning("Renomeio", "Escolha uma pasta de lote válida.")
            return
        if not planilha or not Path(planilha).is_file():
            messagebox.showwarning("Renomeio", "Escolha a planilha.")
            return
        if not tamanhos or not categorias:
            messagebox.showwarning("Renomeio", "Marque pelo menos um tamanho e uma categoria.")
            return
        if not SCRIPT.is_file():
            messagebox.showerror("Renomeio", "Script não encontrado:\n\n" + str(SCRIPT))
            return

        if not simular and not messagebox.askyesno(
            "Renomear",
            "Os arquivos e pastas do lote serão RENOMEADOS de verdade:\n\n"
            + lote + "\n\nRecomendado simular antes. Continuar?",
        ):
            return

        self._salvar_config()

        args = [
            "--lote", lote,
            "--planilha", planilha,
            "--col-nome", self.col_nome.get(),
            "--col-codigo", self.col_codigo.get(),
            "--tamanhos", ",".join(tamanhos),
            "--categorias", ",".join(categorias),
            "--sem-pausa",
            "--simular" if simular else "--real",
        ]
        if self.aba.get():
            args += ["--aba", self.aba.get()]

        self.resumo = None
        self._limpar_saida()
        self._escrever("> PASTA RENOMEIO.py " + " ".join(f'"{a}"' if " " in a else a for a in args) + "\n")
        self._habilitar(False)
        self.status.set("Simulando..." if simular else "Renomeando...")

        ambiente = dict(os.environ)
        ambiente["PYTHONIOENCODING"] = "utf-8"
        ambiente["PYTHONUNBUFFERED"] = "1"

        try:
            self.processo = subprocess.Popen(
                [python_console(), "-u", str(SCRIPT)] + args,
                cwd=str(PASTA),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                stdin=subprocess.DEVNULL,
                env=ambiente,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except Exception as e:
            self.processo = None
            self._habilitar(True)
            self.status.set("Não foi possível iniciar o script.")
            self._escrever(str(e) + "\n")
            return

        threading.Thread(target=self._ler_saida, daemon=True).start()

    def _ler_saida(self):
        processo = self.processo
        for bruto in iter(processo.stdout.readline, b""):
            self.root.after(0, self._escrever, bruto.decode("utf-8", errors="replace"))
        codigo = processo.wait()
        self.root.after(0, self._finalizar, codigo)

    def _finalizar(self, codigo):
        self.processo = None
        self._habilitar(True)
        self._atualizar_pastas()
        if codigo != 0:
            self.status.set(f"Terminou com erro (código {codigo}). Veja a saída.")
        elif self.resumo:
            self.status.set("Concluído: " + self.resumo)
        else:
            self.status.set("Concluído.")

    # --------------------------------------------------------
    # Utilidades
    # --------------------------------------------------------

    def _habilitar(self, habilitar):
        estado = "normal" if habilitar else "disabled"
        self.botao_simular.configure(state=estado)
        self.botao_renomear.configure(state=estado)

    def _limpar_saida(self):
        self.texto_saida.configure(state="normal")
        self.texto_saida.delete("1.0", "end")
        self.texto_saida.configure(state="disabled")

    def _escrever(self, texto):
        m = RE_RESUMO.search(texto)
        if m:
            self.resumo = (
                f"{m.group(1).capitalize()} {m.group(2)} | Já ok {m.group(3)} | "
                f"Não encontrados {m.group(4)} | Erros {m.group(5)}"
            )
        self.texto_saida.configure(state="normal")
        self.texto_saida.insert("end", texto)
        self.texto_saida.see("end")
        self.texto_saida.configure(state="disabled")


def main():
    root = tk.Tk()
    JanelaRenomeio(root)
    root.mainloop()


if __name__ == "__main__":
    main()
